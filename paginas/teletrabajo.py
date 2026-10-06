import calendar
import html
import io
from copy import copy
from datetime import date, timedelta
from pathlib import Path

import openpyxl
import pandas as pd
import streamlit as st
from sqlalchemy import text

import db
from reglas_vacaciones import jefes_por_area

MESES = ["enero", "febrero", "marzo", "abril", "mayo", "junio", "julio",
         "agosto", "septiembre", "octubre", "noviembre", "diciembre"]
DIAS_SEMANA = ["Lun", "Mar", "Mié", "Jue", "Vie", "Sáb", "Dom"]
DIAS_CORTOS = ["LUN", "MAR", "MIÉ", "JUE", "VIE", "SÁB", "DOM"]
DIAS_LARGOS_ES = ["lunes", "martes", "miércoles", "jueves", "viernes", "sábado", "domingo"]
PLANTILLA = Path(__file__).resolve().parent.parent / "plantillas" / "FORMATO_TELETRABAJO.xlsx"
INICIAL = "LMMJVSD"
# Si un feriado cae en uno de estos días (0 = lunes ... 6 = domingo), esa semana
# completa (lunes a domingo) queda sin teletrabajo. Para incluir sábados: {0, 1, 2, 3, 4, 5}
DIAS_QUE_BLOQUEAN = {0, 1, 2, 3, 4}

st.markdown(
    """
<style>
.tt-banda {
  background: #0B4F9C; color: #FFFFFF; border-radius: 14px;
  padding: 1.2rem 1.6rem; display: flex; justify-content: space-between;
  align-items: flex-end; gap: 1rem; flex-wrap: wrap; margin-bottom: 0.8rem;
}
.tt-banda h2 { color: #FFFFFF; margin: 0; padding: 0; font-size: 1.6rem; font-weight: 600; }
.tt-banda p { margin: 0.25rem 0 0; opacity: 0.85; }
.tt-banda .tt-num { text-align: right; line-height: 1; }
.tt-banda .tt-num b { font-size: 2.6rem; display: block; }
.tt-banda .tt-num span { opacity: 0.85; font-size: 0.95rem; }
.tt-tabla { overflow-x: auto; border: 1px solid #D5E2F3; border-radius: 10px; }
.tt-tabla table { border-collapse: collapse; font-size: 0.85rem; width: 100%; }
.tt-tabla th, .tt-tabla td { text-align: center; padding: 4px 2px; min-width: 26px;
  border-bottom: 1px solid #EDF3FB; }
.tt-tabla th { font-weight: 600; color: #14294B; background: #FFFFFF; }
.tt-tabla th span { display: block; font-weight: 400; color: #6A7C99; font-size: 0.75rem; }
.tt-tabla .nom { text-align: left; white-space: nowrap; padding: 4px 12px;
  position: sticky; left: 0; background: #FFFFFF; min-width: 200px; }
.tt-tabla .fin { background: #F4F7FB; }
.tt-tabla .hoy { box-shadow: inset 0 0 0 2px #0B4F9C; }
.tt-tabla td.on.pend i { background: #9DBDE6; }
.tt-chips span.pend { background: #FFFFFF; border: 1px dashed #9DBDE6; }
.tt-tabla td.on i { display: inline-block; width: 16px; height: 16px;
  border-radius: 4px; background: #0B4F9C; }
.tt-tabla tfoot td { color: #4A5D7E; font-weight: 600; border-bottom: none; }
.tt-tabla .fer { background: repeating-linear-gradient(135deg, #FDECEC 0 4px, #FFFFFF 4px 8px); }
.tt-tabla .sem { background: #FEF6F5; }
.tt-tabla th.fer s { color: #B42318; }
.tt-leyenda { font-size: 0.85rem; color: #4A5D7E; margin: 0.4rem 0 1rem; }
.tt-leyenda em { display: inline-block; width: 12px; height: 12px; border-radius: 3px;
  margin: 0 4px 0 12px; vertical-align: -1px; }
.tt-cal-cab { text-align: center; font-weight: 600; color: #4A5D7E; }
.tt-chips span.fer { background: #FDECEC; color: #B42318; }
.tt-chips span { display: inline-block; background: #EDF3FB; color: #0B4F9C;
  border-radius: 999px; padding: 0.3rem 0.8rem; margin: 0 0.4rem 0.4rem 0; font-weight: 500; }
</style>
""",
    unsafe_allow_html=True,
)


MENORES = {"y", "de", "del", "la", "las", "los", "e"}


def pascua(anio):
    a, b, c = anio % 19, anio // 100, anio % 100
    d, e = b // 4, b % 4
    f = (b + 8) // 25
    g = (b - f + 1) // 3
    h = (19 * a + b - d - g + 15) % 30
    i, k = c // 4, c % 4
    l = (32 + 2 * e + 2 * i - h - k) % 7
    m = (a + 11 * h + 22 * l) // 451
    mes = (h + l - 7 * m + 114) // 31
    dia = (h + l - 7 * m + 114) % 31 + 1
    return date(anio, mes, dia)


def feriados_peru(anio):
    """Feriados nacionales de Perú (los días no laborables por decreto se marcan a mano)."""
    p = pascua(anio)
    fijos = {
        (1, 1): "Año Nuevo", (5, 1): "Día del Trabajo",
        (6, 7): "Batalla de Arica y Día de la Bandera", (6, 29): "San Pedro y San Pablo",
        (7, 23): "Día de la Fuerza Aérea del Perú", (7, 28): "Fiestas Patrias",
        (7, 29): "Fiestas Patrias", (8, 6): "Batalla de Junín", (8, 30): "Santa Rosa de Lima",
        (10, 8): "Combate de Angamos", (11, 1): "Día de Todos los Santos",
        (12, 8): "Inmaculada Concepción", (12, 9): "Batalla de Ayacucho", (12, 25): "Navidad",
    }
    lista = {date(anio, m, d).isoformat(): n for (m, d), n in fijos.items()}
    lista[(p - timedelta(days=3)).isoformat()] = "Jueves Santo"
    lista[(p - timedelta(days=2)).isoformat()] = "Viernes Santo"
    return lista


def lunes_de(d):
    return d - timedelta(days=d.weekday())


def semanas_bloqueadas(feriados):
    """Lunes de cada semana que tiene un feriado en día laborable -> nombres de los feriados."""
    semanas = {}
    for iso, nombre in sorted(feriados.items()):
        f = date.fromisoformat(iso)
        if f.weekday() in DIAS_QUE_BLOQUEAN:
            semanas.setdefault(lunes_de(f), []).append(nombre or "Feriado")
    return semanas


def cargar_feriados():
    df = db.consultar("SELECT fecha, nombre FROM feriados")
    return dict(zip(df["fecha"], df["nombre"].fillna("")))


def alternar_feriado(iso, es_feriado):
    if es_feriado:
        db.ejecutar("DELETE FROM feriados WHERE fecha = :f", {"f": iso})
        st.session_state["tt_aviso"] = f"Se quitó el feriado del {date.fromisoformat(iso):%d/%m/%Y}."
    else:
        nombre = (st.session_state.get("tt_fer_nombre") or "").strip() or "Feriado"
        db.ejecutar("INSERT INTO feriados (fecha, nombre) VALUES (:f, :n)", {"f": iso, "n": nombre})
        st.session_state["tt_fer_nombre"] = ""
        st.session_state["tt_aviso"] = f"{date.fromisoformat(iso):%d/%m/%Y} marcado como feriado ({nombre})."


def bonito(nombre):
    palabras = str(nombre).title().split(" ")
    return " ".join(w.lower() if i and w.lower() in MENORES else w for i, w in enumerate(palabras))


def cargar():
    personal = db.consultar("SELECT cod, nombre, area FROM rrhh_personal ORDER BY nombre")
    todos = db.consultar(
        """SELECT t.id, t.cod, t.fecha, t.estado, t.resuelto_por, p.nombre, p.area, p.puesto,
                  p.departamento, p.unidad, p.sede
           FROM teletrabajo_solicitudes t JOIN rrhh_personal p ON p.cod = t.cod"""
    )
    return personal, todos[todos["estado"] != "RECHAZADO"], todos[todos["estado"] == "RECHAZADO"]


def dias_del_mes(anio, mes):
    return [date(anio, mes, d) for d in range(1, calendar.monthrange(anio, mes)[1] + 1)]


def tabla_mes(tele_mes, anio, mes, hoy, feriados):
    dias = dias_del_mes(anio, mes)
    marcados = dict(zip(zip(tele_mes["cod"], tele_mes["fecha"]), tele_mes["estado"]))
    personas = tele_mes[["cod", "nombre"]].drop_duplicates().sort_values("nombre")

    def clase(d):
        c = []
        if d.isoformat() in feriados:
            c.append("fer")
        elif lunes_de(d) in bloqueadas:
            c.append("sem")
        elif d.weekday() >= 5:
            c.append("fin")
        if d == hoy:
            c.append("hoy")
        return " ".join(c)

    def numero(d):
        iso = d.isoformat()
        if iso in feriados:
            return f'<s title="{html.escape(feriados[iso] or "Feriado")}">{d.day}</s>'
        return str(d.day)

    cab = "".join(f'<th class="{clase(d)}"><span>{INICIAL[d.weekday()]}</span>{numero(d)}</th>' for d in dias)
    filas = []
    for r in personas.itertuples():
        celdas, total = [], 0
        for d in dias:
            estado = marcados.get((r.cod, d.isoformat()))
            on = estado is not None
            total += on
            extra = (" on pend" if estado == "PENDIENTE" else " on") if on else ""
            titulo = ' title="En espera de aprobación"' if estado == "PENDIENTE" else ""
            celdas.append(f'<td class="{clase(d)}{extra}"{titulo}>{"<i></i>" if on else ""}</td>')
        filas.append(f'<tr><td class="nom">{html.escape(bonito(r.nombre))}</td>{"".join(celdas)}<td><b>{total}</b></td></tr>')
    por_dia = tele_mes.groupby("fecha").size()
    pie = "".join(f'<td class="{clase(d)}">{por_dia.get(d.isoformat(), "") or ""}</td>' for d in dias)
    st.markdown(
        f"""<div class="tt-tabla"><table>
<thead><tr><th class="nom"></th>{cab}<th>Total</th></tr></thead>
<tbody>{"".join(filas)}</tbody>
<tfoot><tr><td class="nom">Personas por día</td>{pie}<td>{len(tele_mes)}</td></tr></tfoot>
</table></div>
<div class="tt-leyenda"><em style="background:#0B4F9C"></em>teletrabajo aprobado
<em style="background:#9DBDE6"></em>en espera de aprobación del jefe
<em style="background:repeating-linear-gradient(135deg,#FDECEC 0 3px,#FFFFFF 3px 6px);border:1px solid #F3B8B5"></em>feriado
<em style="background:#FEF6F5;border:1px solid #F3B8B5"></em>semana de feriado (sin teletrabajo)
<em style="background:#F4F7FB;border:1px solid #D5E2F3"></em>fin de semana</div>""",
        unsafe_allow_html=True,
    )


def semanas_del_mes(anio, mes):
    """Semanas (lunes a domingo) que tienen días hábiles del mes: [(lunes, [días hábiles])]. Máximo 5."""
    semanas = []
    for semana in calendar.Calendar(firstweekday=0).monthdatescalendar(anio, mes):
        habiles = [d for d in semana if d.month == mes and d.weekday() < 5]
        if habiles:
            semanas.append((semana[0], habiles))
    return semanas[:5]


def excel_formato(vista, anio, mes):
    """Llena FORMATO_TELETRABAJO.xlsx: una fila por persona y sus días por semana."""
    wb = openpyxl.load_workbook(PLANTILLA)
    ws = wb.active
    ws["H1"] = f"{MESES[mes - 1].upper()} {anio}"
    semanas = semanas_del_mes(anio, mes)
    for i in range(5):
        celda = ws.cell(row=3, column=8 + i)
        if i < len(semanas):
            lun, habiles = semanas[i]
            rango = f"{habiles[0]:%d/%m} al {habiles[-1]:%d/%m}"
            celda.value = f"{rango}\nFERIADO" if lun in bloqueadas else rango
        else:
            celda.value = "—"

    dias_por_persona = {}
    for cod, f in zip(vista["cod"], vista["fecha"]):
        dias_por_persona.setdefault(cod, []).append(date.fromisoformat(f))
    personas = vista.drop_duplicates("cod").sort_values("nombre")

    for fila, p in enumerate(personas.itertuples(), start=4):
        if fila > 33:  # la plantilla trae 30 filas con formato; si hay más, se copia el estilo
            for col in range(1, 13):
                origen, destino = ws.cell(row=33, column=col), ws.cell(row=fila, column=col)
                destino._style = copy(origen._style)
        for col, valor in enumerate([p.cod, p.nombre, p.puesto, p.area, p.departamento, p.unidad, p.sede], start=1):
            celda = ws.cell(row=fila, column=col, value=valor)
            al = copy(celda.alignment)
            al.vertical = "center"
            celda.alignment = al
        for i, (lun, _) in enumerate(semanas):
            celda = ws.cell(row=fila, column=8 + i)
            if lun in bloqueadas:
                celda.value = "FERIADO"
            else:
                dias = sorted(d for d in dias_por_persona[p.cod] if lunes_de(d) == lun)
                celda.value = "\n".join(f"{DIAS_CORTOS[d.weekday()]} {d:%d/%m}" for d in dias) or None
            al = copy(celda.alignment)
            al.wrap_text = True
            al.vertical = "center"
            celda.alignment = al

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def banda(titulo, subtitulo, numero, etiqueta):
    st.markdown(
        f"""<div class="tt-banda">
  <div><h2>{html.escape(titulo)}</h2><p>{html.escape(subtitulo)}</p></div>
  <div class="tt-num"><b>{numero}</b><span>{html.escape(etiqueta)}</span></div>
</div>""",
        unsafe_allow_html=True,
    )


# ---------- Datos y sesión ----------
sesion = st.session_state.get("sesion", {})
rol = sesion.get("rol")
mi_cod = sesion.get("cod")
hoy = db.hoy()
personal, tele, rechazados = cargar()
jefes = jefes_por_area(personal)          # {área: cod del jefe}
cods_jefes = set(jefes.values())


def lo_aprueba_el_admin(cod, area):
    """Jefes y áreas sin jefe: los aprueba el admin, así que lo que él programa queda aprobado."""
    return cod in cods_jefes or area not in jefes
feriados = cargar_feriados()
bloqueadas = semanas_bloqueadas(feriados)

# ---------- Selector de mes ----------
if "tt_mes" not in st.session_state:
    st.session_state["tt_mes"] = hoy.replace(day=1)
mes_ini = st.session_state["tt_mes"]

st.title("Teletrabajo")
n1, n2, n3, _ = st.columns([1, 1, 1, 3])
if n1.button("Mes anterior"):
    st.session_state["tt_mes"] = (mes_ini - timedelta(days=1)).replace(day=1)
    st.rerun()
if n2.button("Este mes", disabled=mes_ini == hoy.replace(day=1)):
    st.session_state["tt_mes"] = hoy.replace(day=1)
    st.rerun()
if n3.button("Mes siguiente"):
    st.session_state["tt_mes"] = (mes_ini + timedelta(days=32)).replace(day=1)
    st.rerun()

anio, mes = mes_ini.year, mes_ini.month
nombre_mes = f"{MESES[mes - 1].capitalize()} {anio}"
tele_mes = tele[tele["fecha"].str.startswith(f"{anio}-{mes:02d}", na=False)]

if "tt_aviso" in st.session_state:
    st.success(st.session_state.pop("tt_aviso"))

# ---------- Colaborador: solo sus días ----------
if rol == "colaborador":
    mios = tele_mes[tele_mes["cod"] == mi_cod].sort_values("fecha")
    banda(nombre_mes, "Tus días de teletrabajo", len(mios), "días este mes")
    if not mi_cod:
        st.warning("Tu usuario todavía no está asociado a tu ficha de personal. Avisa al administrador.")
    elif mios.empty:
        st.info("No tienes días de teletrabajo en este mes.")
    else:
        chips = "".join(
            f'<span class="pend">{DIAS_SEMANA[d.weekday()]} {d:%d/%m} (en espera)</span>' if e == "PENDIENTE"
            else f"<span>{DIAS_SEMANA[d.weekday()]} {d:%d/%m}</span>"
            for d, e in zip(pd.to_datetime(mios["fecha"]).dt.date, mios["estado"])
        )
        st.markdown(f'<div class="tt-chips">{chips}</div>', unsafe_allow_html=True)
    fer_mes = sorted(f for f in feriados if f.startswith(f"{anio}-{mes:02d}"))
    if fer_mes:
        st.markdown("**Feriados del mes**")
        st.markdown('<div class="tt-chips">' + "".join(
            f'<span class="fer">{date.fromisoformat(f):%d/%m} {html.escape(feriados[f])}</span>' for f in fer_mes
        ) + "</div>", unsafe_allow_html=True)
    st.caption("El teletrabajo lo programa el administrador y lo aprueba tu jefe. Si necesitas un cambio, coordínalo con ellos.")
    st.stop()

# ---------- Jefe: su área, solo lectura ----------
if rol == "jefe":
    yo = personal[personal["cod"] == mi_cod]
    if yo.empty:
        st.warning("Tu usuario no está asociado a tu ficha de personal. Avisa al administrador.")
        st.stop()
    mis_areas = [a for a, c in jefes.items() if c == mi_cod] or [yo.iloc[0]["area"]]
    for area in mis_areas:
        del_area = tele_mes[(tele_mes["area"] == area) & (tele_mes["cod"] != mi_cod)]
        banda(nombre_mes, f"Teletrabajo en {bonito(area)}", del_area["cod"].nunique(), "personas este mes")
        if del_area.empty:
            st.info("Nadie de tu área tiene teletrabajo en este mes.")
        else:
            tabla_mes(del_area, anio, mes, hoy, feriados)
    st.caption("El teletrabajo lo programa el administrador. Los días en espera los apruebas en «Aprobaciones».")
    st.stop()

# ---------- Admin ----------
areas = ["Todas las áreas"] + sorted(personal["area"].dropna().unique().tolist())
banda(nombre_mes, "Teletrabajo del equipo", tele_mes["cod"].nunique(), f"personas, {len(tele_mes)} días en total")

tab_cal, tab_prog, tab_quitar, tab_fer = st.tabs(["Calendario", "Programar días", "Quitar días", "Feriados"])

with tab_cal:
    area_cal = st.selectbox("Área", areas, key="tt_area_cal")
    vista = tele_mes if area_cal == areas[0] else tele_mes[tele_mes["area"] == area_cal]
    if vista.empty:
        st.info("No hay teletrabajo registrado en este mes. Programa días en la pestaña «Programar días».")
    else:
        tabla_mes(vista, anio, mes, hoy, feriados)
        semanas_p = vista.assign(lunes=pd.to_datetime(vista["fecha"]).dt.date.map(lunes_de))
        dobles = semanas_p.groupby(["nombre", "lunes"]).size()
        dobles = dobles[dobles > 1]
        if not dobles.empty:
            detalle = "; ".join(f"{bonito(n)} ({c} días la semana del {l:%d/%m})" for (n, l), c in dobles.items())
            st.warning(f"Solo se permite un día de teletrabajo por semana. Revisa: {detalle}. "
                       "Puedes corregirlo en «Quitar días».")
        aprobados = vista[vista["estado"] == "APROBADO"]
        en_espera = len(vista) - len(aprobados)
        if en_espera:
            st.info(f"{en_espera} {'día está' if en_espera == 1 else 'días están'} en espera de aprobación "
                    "de los jefes. El Excel solo incluye los días aprobados.")
        if PLANTILLA.exists() and aprobados.empty:
            st.caption("Todavía no hay días aprobados para exportar en este mes.")
        elif PLANTILLA.exists():
            st.download_button(
                "Descargar formato de programación (Excel)",
                data=excel_formato(aprobados, anio, mes),
                file_name=f"PROGRAMACION_TELETRABAJO_{MESES[mes - 1].upper()}_{anio}.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                type="primary",
            )
        else:
            st.warning("Falta la plantilla plantillas/FORMATO_TELETRABAJO.xlsx para exportar.")
    rech_mes = rechazados[rechazados["fecha"].str.startswith(f"{anio}-{mes:02d}", na=False)]
    if area_cal != areas[0]:
        rech_mes = rech_mes[rech_mes["area"] == area_cal]
    if not rech_mes.empty:
        with st.expander(f"Rechazados por los jefes este mes ({len(rech_mes)})"):
            for r in rech_mes.sort_values(["nombre", "fecha"]).itertuples():
                d = date.fromisoformat(r.fecha)
                st.markdown(f"{bonito(r.nombre)}: {DIAS_SEMANA[d.weekday()]} {d:%d/%m}"
                            + (f", rechazado por {bonito(r.resuelto_por)}" if r.resuelto_por else ""))
            st.caption("Puedes volver a programar esas fechas u otras en «Programar días».")

with tab_prog:
    area_p = st.selectbox("Filtrar personas por área", areas, key="tt_area_prog")
    lista = personal if area_p == areas[0] else personal[personal["area"] == area_p]
    etiquetas = {f"{bonito(r.nombre)} ({r.cod})": r.cod for r in lista.itertuples()}
    elegidos = st.multiselect("Personas", list(etiquetas.keys()), placeholder="Elige una o varias personas")

    c1, c2 = st.columns(2)
    desde = c1.date_input("Desde", value=max(hoy, mes_ini), format="DD/MM/YYYY", key="tt_desde")
    hasta = c2.date_input("Hasta", value=desde + timedelta(days=27), format="DD/MM/YYYY", key="tt_hasta")
    c3, c4 = st.columns(2)
    habiles = DIAS_SEMANA[:5]
    dia_a = habiles.index(c3.selectbox("Una semana", habiles, index=1, key="tt_dia_a"))
    dia_b = habiles.index(c4.selectbox("La semana siguiente", habiles, index=3, key="tt_dia_b"))
    un_dia = desde == hasta
    if un_dia:
        st.caption("Misma fecha en Desde y Hasta: se registra solo ese día (si esa semana no tiene otro).")
    else:
        st.caption(f"Un solo día de teletrabajo por semana, intercalado: una semana {DIAS_LARGOS_ES[dia_a]}, "
                   f"la siguiente {DIAS_LARGOS_ES[dia_b]}. Si la persona ya tenía teletrabajo antes, "
                   "se continúa su alternancia.")

    def describir_bloqueos(dias):
        partes = []
        for lun in sorted({lunes_de(d) for d in dias if lunes_de(d) in bloqueadas}):
            partes.append(f"semana del {lun:%d/%m} ({', '.join(bloqueadas[lun])})")
        sueltos = [d for d in dias if lunes_de(d) not in bloqueadas]
        partes += [f"{d:%d/%m} ({feriados[d.isoformat()]})" for d in sueltos]
        return "; ".join(partes)

    def plan_persona(cod):
        """Fechas propuestas para una persona: una por semana, intercalando los dos días."""
        if un_dia:
            return [desde]
        previos = tele[(tele["cod"] == cod) & (tele["fecha"] < desde.isoformat())]
        ultimo = date.fromisoformat(previos["fecha"].max()) if not previos.empty else None
        if ultimo and ultimo.weekday() in (dia_a, dia_b) and dia_a != dia_b:
            base_lun, base_dia = lunes_de(ultimo), ultimo.weekday()   # continúa la alternancia
        else:
            base_lun, base_dia = lunes_de(desde) - timedelta(days=7), dia_b  # la 1.ª semana toca el primer día
        otro = {dia_a: dia_b, dia_b: dia_a}
        fechas_p, lun = [], lunes_de(desde)
        while lun <= hasta:
            k = (lun - base_lun).days // 7
            dia = base_dia if k % 2 == 0 else otro.get(base_dia, dia_a)
            f = lun + timedelta(days=dia)
            if desde <= f <= hasta:
                fechas_p.append(f)
            lun += timedelta(days=7)
        return fechas_p

    cods = [etiquetas[e] for e in elegidos]
    nuevos, en_feriado, ya_semana, en_vacaciones, vista_previa = [], [], 0, [], []
    if hasta >= desde and dia_a != dia_b and cods:
        vac = db.consultar(
            "SELECT cod, inicio, fin FROM vacaciones_solicitudes WHERE estado <> 'RECHAZADA'"
        )
        nombres = {v: k.rsplit(" (", 1)[0] for k, v in etiquetas.items()}
        for cod in cods:
            vac_p = vac[vac["cod"] == cod]
            semanas_con = {lunes_de(date.fromisoformat(f)) for f in tele[tele["cod"] == cod]["fecha"]}
            suyas = []
            for f in plan_persona(cod):
                iso = f.isoformat()
                if iso in feriados or lunes_de(f) in bloqueadas:
                    en_feriado.append(f)
                elif lunes_de(f) in semanas_con:
                    ya_semana += 1
                elif ((vac_p["inicio"] <= iso) & (vac_p["fin"] >= iso)).any():
                    en_vacaciones.append((cod, f))
                else:
                    nuevos.append({"cod": cod, "fecha": iso})
                    semanas_con.add(lunes_de(f))
                    suyas.append(f)
            vista_previa.append({
                "Colaborador": nombres[cod],
                "Días que se registran": ", ".join(f"{DIAS_SEMANA[d.weekday()]} {d:%d/%m}" for d in suyas) or "Ninguno",
            })

    if hasta < desde:
        st.error("La fecha «Hasta» es anterior a «Desde».")
    elif dia_a == dia_b and not un_dia:
        st.error("Elige dos días distintos: el teletrabajo se intercala entre una semana y la siguiente.")
    elif cods:
        st.markdown(
            f"Vas a registrar **{len(nuevos)} {'día' if len(nuevos) == 1 else 'días'}** de teletrabajo "
            f"para **{len(cods)} {'persona' if len(cods) == 1 else 'personas'}**."
        )
        if vista_previa:
            st.dataframe(pd.DataFrame(vista_previa), width="stretch", hide_index=True)
        if en_feriado:
            unicos = sorted(set(en_feriado))
            st.info(f"Se omiten las semanas con feriado: {describir_bloqueos(unicos)}.")
        if ya_semana:
            st.info(f"{ya_semana} {'semana ya tenía' if ya_semana == 1 else 'semanas ya tenían'} teletrabajo "
                    "registrado; se dejan como están (máximo un día por semana).")
        if en_vacaciones:
            detalle = ", ".join(f"{nombres[c]} el {f:%d/%m}" for c, f in en_vacaciones[:6])
            extra = f" y {len(en_vacaciones) - 6} más" if len(en_vacaciones) > 6 else ""
            st.warning(f"Se omiten {len(en_vacaciones)} fechas porque coinciden con vacaciones: {detalle}{extra}.")

    area_de = dict(zip(personal["cod"], personal["area"]))
    if nuevos:
        directos = sum(lo_aprueba_el_admin(n["cod"], area_de.get(n["cod"])) for n in nuevos)
        if directos < len(nuevos):
            st.caption("Los días quedan en espera hasta que el jefe del área los apruebe en «Aprobaciones». "
                       "Los de jefes y áreas sin jefe quedan aprobados directamente, porque esos los apruebas tú.")

    if st.button("Registrar teletrabajo", type="primary", disabled=not nuevos):
        quien = sesion.get("usuario")
        ahora = db.ahora()
        for n in nuevos:
            directo = lo_aprueba_el_admin(n["cod"], area_de.get(n["cod"]))
            n.update(creado=ahora, estado="APROBADO" if directo else "PENDIENTE", quien=quien if directo else None)
        with db.motor().begin() as con:
            con.execute(
                text("INSERT INTO teletrabajo_solicitudes (cod, fecha, estado, creado, resuelto_por) "
                     "VALUES (:cod, :fecha, :estado, :creado, :quien)"),
                nuevos,
            )
        espera = sum(n["estado"] == "PENDIENTE" for n in nuevos)
        aviso = f"Listo: se {'registró' if len(nuevos) == 1 else 'registraron'} {len(nuevos)} {'día' if len(nuevos) == 1 else 'días'} de teletrabajo."
        if espera:
            aviso += f" {espera} {'queda' if espera == 1 else 'quedan'} en espera de aprobación del jefe."
        st.session_state["tt_aviso"] = aviso
        st.rerun()

with tab_quitar:
    con_dias = tele_mes[["cod", "nombre"]].drop_duplicates().sort_values("nombre")
    if con_dias.empty:
        st.info(f"No hay días registrados en {nombre_mes.lower()}.")
    else:
        opciones = {f"{bonito(r.nombre)} ({r.cod})": r.cod for r in con_dias.itertuples()}
        persona = st.selectbox("Persona", list(opciones.keys()), key="tt_quitar_p")
        cod_q = opciones[persona]
        suyos = tele_mes[tele_mes["cod"] == cod_q].sort_values("fecha")
        etiq_f = {
            f"{DIAS_SEMANA[date.fromisoformat(f).weekday()]} {date.fromisoformat(f):%d/%m/%Y}": int(i)
            for i, f in zip(suyos["id"], suyos["fecha"])
        }
        quitar = st.multiselect("Días a quitar", list(etiq_f.keys()), placeholder="Elige los días")
        if st.button("Quitar días seleccionados", disabled=not quitar):
            with db.motor().begin() as con:
                con.execute(text("DELETE FROM teletrabajo_solicitudes WHERE id = :id"),
                            [{"id": etiq_f[q]} for q in quitar])
            st.session_state["tt_aviso"] = f"Listo: se {'quitó' if len(quitar) == 1 else 'quitaron'} {len(quitar)} {'día' if len(quitar) == 1 else 'días'} de {persona.rsplit(' (', 1)[0]}."
            st.rerun()

with tab_fer:
    st.markdown(f"**{nombre_mes}.** Toca un día para marcarlo como feriado; tócalo otra vez para quitarlo. "
                "Los feriados aparecen tachados, y en la semana de un feriado no se puede programar teletrabajo.")
    st.text_input("Nombre del feriado (opcional, antes de tocar el día)", key="tt_fer_nombre",
                  placeholder="Ej.: Día no laborable")
    cab = st.columns(7)
    for i, nombre_d in enumerate(DIAS_SEMANA):
        cab[i].markdown(f'<div class="tt-cal-cab">{nombre_d}</div>', unsafe_allow_html=True)
    for semana in calendar.Calendar(firstweekday=0).monthdatescalendar(anio, mes):
        cols = st.columns(7)
        for i, d in enumerate(semana):
            if d.month != mes:
                continue
            iso = d.isoformat()
            es_fer = iso in feriados
            cols[i].button(
                f"~~{d.day}~~" if es_fer else str(d.day),
                key=f"fer_{iso}",
                type="primary" if es_fer else "secondary",
                width="stretch",
                help=(feriados[iso] or "Feriado") if es_fer else None,
                on_click=alternar_feriado,
                args=(iso, es_fer),
            )

    fer_mes = sorted(f for f in feriados if f.startswith(f"{anio}-{mes:02d}"))
    if fer_mes:
        st.markdown('<div class="tt-chips">' + "".join(
            f'<span class="fer">{date.fromisoformat(f):%d/%m} {html.escape(feriados[f])}</span>' for f in fer_mes
        ) + "</div>", unsafe_allow_html=True)

    sem_mes = sorted(l for l in bloqueadas
                     if l.replace(day=1) <= mes_ini <= (l + timedelta(days=6)).replace(day=1))
    if sem_mes:
        st.caption("Semanas sin teletrabajo: " + "; ".join(
            f"del {l:%d/%m} al {l + timedelta(days=6):%d/%m}" for l in sem_mes) + ".")
    fechas_tele = pd.to_datetime(tele["fecha"]).dt.date
    choque = tele[fechas_tele.map(lambda d: lunes_de(d) in sem_mes or d.isoformat() in fer_mes)]
    if not choque.empty:
        uno = len(choque) == 1
        st.warning(f"Hay {len(choque)} {'día' if uno else 'días'} de teletrabajo "
                   f"{'registrado' if uno else 'registrados'} en feriados o en semanas con feriado.")
        if st.button("Quitar ese teletrabajo"):
            with db.motor().begin() as con:
                con.execute(text("DELETE FROM teletrabajo_solicitudes WHERE id = :id"),
                            [{"id": int(i)} for i in choque["id"]])
            st.session_state["tt_aviso"] = "Listo: se quitó el teletrabajo de las semanas con feriado."
            st.rerun()

    st.divider()
    nacionales = feriados_peru(anio)
    faltan = {f: n for f, n in nacionales.items() if f not in feriados}
    st.markdown(f"**Feriados nacionales de {anio}**")
    if faltan:
        st.caption(f"Faltan {len(faltan)} de los {len(nacionales)} feriados nacionales. "
                   "Los días no laborables que dicte el Gobierno márcalos a mano.")
        if st.button(f"Marcar los feriados nacionales de {anio}"):
            with db.motor().begin() as con:
                con.execute(text("INSERT INTO feriados (fecha, nombre) VALUES (:f, :n) ON CONFLICT (fecha) DO NOTHING"),
                            [{"f": f, "n": n} for f, n in faltan.items()])
            st.session_state["tt_aviso"] = f"Listo: se marcaron {len(faltan)} feriados nacionales de {anio}."
            st.rerun()
    else:
        st.caption(f"Los {len(nacionales)} feriados nacionales de {anio} ya están marcados.")
