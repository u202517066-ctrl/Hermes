import calendar
import html
from datetime import date, timedelta

import pandas as pd
import streamlit as st

import correo
import db
import estilo
from reglas_vacaciones import (
    APROBADORES, DIAS_BLOQUE, MESES_CICLO,
    inicio_ciclo, jefes_por_area, primer_aniversario, sumar_meses, usados_en_ciclo,
)

st.markdown(
    """
<style>
.ap-banda {
  background: #0B4F9C; color: #FFFFFF; border-radius: 14px;
  padding: 1.2rem 1.6rem; display: flex; justify-content: space-between;
  align-items: flex-end; gap: 1rem; flex-wrap: wrap; margin: 0.6rem 0 0.8rem;
}
.ap-banda h2 { color: #FFFFFF; margin: 0; padding: 0; font-size: 1.45rem; font-weight: 600; }
.ap-banda p { margin: 0.25rem 0 0; opacity: 0.85; }
.ap-num { text-align: right; line-height: 1; }
.ap-num b { font-size: 2.6rem; display: block; }
.ap-num span { opacity: 0.85; font-size: 0.95rem; }
.a-cab { display: flex; justify-content: space-between; align-items: baseline; margin: 6px 0 10px; }
.a-cab b { font-size: 1.05rem; color: var(--h-tinta); }
.a-cab span { color: var(--h-suave); font-size: 0.9rem; }
.a-quien h4 { margin: 0; padding: 0; font-size: 1rem; color: var(--h-tinta); }
.a-quien p { margin: 2px 0 0; color: var(--h-suave); font-size: 0.9rem; }
.a-ctx { margin-top: 8px; display: flex; gap: 8px; flex-wrap: wrap; }
.a-pedido { display: grid; grid-template-columns: auto 1fr; gap: 14px; align-items: start; }
[class*="st-key-pedu_"] { border-left: 4px solid var(--h-ambar) !important; }
[class*="st-key-ped"] { background: var(--h-superficie); }
.a-equipo, .a-equipo * { border: 0 !important; }
.a-equipo { width: 100%; border-collapse: collapse !important; margin: 0; }
.a-equipo tr { background: transparent !important; border-bottom: 1px solid var(--h-linea) !important; }
.a-equipo tbody tr:last-child { border-bottom: 0 !important; }
.a-equipo th { text-align: left; font-size: 0.8rem; color: var(--h-suave); font-weight: 600; padding: 6px 8px; }
.a-equipo td { padding: 9px 8px; font-size: 0.9rem; color: var(--h-tinta); vertical-align: middle; }
.a-equipo td small { display: block; color: var(--h-suave); font-size: 0.78rem; }
.a-barra { height: 8px; background: var(--h-fondo); border-radius: 99px; overflow: hidden; min-width: 80px; margin-bottom: 3px; }
.a-barra i { display: block; height: 100%; background: var(--h-azul); border-radius: 99px; }
</style>
""",
    unsafe_allow_html=True,
)

MENORES = {"y", "de", "del", "la", "las", "los", "e"}
ESTADOS = {
    "PENDIENTE": "En espera",
    "PENDIENTE_GERENCIA": "En espera de gerencia",
    "APROBADA": "Aprobada",
    "RECHAZADA": "Rechazada",
}


def bonito(texto):
    palabras = str(texto).title().split(" ")
    return " ".join(w.lower() if i and w.lower() in MENORES else w for i, w in enumerate(palabras))


def fecha(iso, con_anio=True):
    d = date.fromisoformat(str(iso))
    return f"{d:%d/%m/%Y}" if con_anio else f"{d:%d/%m}"


def banda(titulo, subtitulo, numero, etiqueta):
    st.markdown(
        f"""<div class="ap-banda">
  <div><h2>{html.escape(titulo)}</h2><p>{html.escape(subtitulo)}</p></div>
  <div class="ap-num"><b>{numero}</b><span>{html.escape(etiqueta)}</span></div>
</div>""",
        unsafe_allow_html=True,
    )


# ---------- Sesión y datos ----------
st.title("Aprobaciones")
sesion = st.session_state.get("sesion", {})
rol = sesion.get("rol")
mi_cod = sesion.get("cod")
if rol not in ("jefe", "admin"):
    st.error("Solo jefaturas y el administrador pueden ver esta página.")
    st.stop()

hoy = db.hoy()
personal = db.consultar("SELECT cod, nombre, puesto, area, seccion, f_ingreso FROM rrhh_personal ORDER BY nombre")
sols = db.consultar(
    "SELECT id, cod, inicio, fin, dias, bloque, estado, resuelto_por FROM vacaciones_solicitudes ORDER BY inicio"
).merge(personal[["cod", "nombre", "puesto", "area"]], on="cod", how="inner")
activas = sols[sols["estado"] != "RECHAZADA"]
tele = db.consultar("SELECT id, cod, fecha, estado, creado FROM teletrabajo_solicitudes ORDER BY fecha").merge(
    personal[["cod", "nombre", "puesto", "area"]], on="cod", how="inner")
tele_pend = tele[tele["estado"] == "PENDIENTE"]
DIAS_SEMANA = ["Lun", "Mar", "Mié", "Jue", "Vie", "Sáb", "Dom"]

jefes = jefes_por_area(personal)          # {área: cod del jefe}
cods_jefes = set(jefes.values())
mis_areas = [a for a, c in jefes.items() if c == mi_cod]
yo = personal[personal["cod"] == mi_cod]
quien = yo.iloc[0]["nombre"] if not yo.empty else sesion.get("usuario")

if "ap_aviso" in st.session_state:
    st.success(st.session_state.pop("ap_aviso"))


def resolver(id_, nuevo_estado, nombre):
    db.ejecutar(
        "UPDATE vacaciones_solicitudes SET estado = :e, resuelto_por = :q WHERE id = :id",
        {"e": nuevo_estado, "q": quien, "id": int(id_)},
    )
    st.session_state["ap_aviso"] = f"Solicitud de {bonito(nombre)} {ESTADOS[nuevo_estado].lower()}."


def rango(a, b):
    a, b = date.fromisoformat(str(a)), date.fromisoformat(str(b))
    if (a.year, a.month) == (b.year, b.month):
        return f"del {a.day} al {b.day} de {estilo.MESES[b.month - 1]}"
    return f"del {a.day} de {estilo.MESES[a.month - 1]} al {b.day} de {estilo.MESES[b.month - 1]}"


def saldo_num(cod):
    """Días por programar de una persona (contando lo aprobado y lo en espera); None si aún no tiene periodo."""
    fila = personal[personal["cod"] == cod]
    if fila.empty:
        return None
    try:
        ingreso = date.fromisoformat(str(fila.iloc[0]["f_ingreso"]))
    except ValueError:
        return None
    c = inicio_ciclo(ingreso, hoy)
    if c is None:
        return None
    u = usados_en_ciclo(activas, cod, c, sumar_meses(c, MESES_CICLO))
    return 2 * DIAS_BLOQUE - u[1] - u[2]


def companeros(cod):
    """Personas con las que se compara: su grupo de retail, o su área."""
    fila = personal[personal["cod"] == cod]
    if fila.empty:
        return set()
    f = fila.iloc[0]
    if str(f["seccion"] or "").upper().startswith("GRUPO"):
        return set(personal[personal["seccion"] == f["seccion"]]["cod"]) - {cod}
    return set(personal[personal["area"] == f["area"]]["cod"]) - {cod}


def tarjetas(df, prefijo):
    if df.empty:
        st.markdown('<div class="h-panel" style="color:var(--h-suave)">No hay solicitudes por aprobar.</div>',
                    unsafe_allow_html=True)
        return
    for r in df.sort_values("inicio").itertuples():
        ini_d = date.fromisoformat(r.inicio)
        faltan = (ini_d - hoy).days
        cuando = ("Empieza hoy" if faltan == 0 else "Empieza mañana" if faltan == 1 else
                  f"Empieza en {faltan} días" if faltan > 1 else f"Empezó hace {-faltan} {'día' if faltan == -1 else 'días'}")
        cruces = activas[activas["cod"].isin(companeros(r.cod)) & (activas["id"] != r.id)
                         & (activas["inicio"] <= r.fin) & (activas["fin"] >= r.inicio)]
        ctx = []
        if cruces.empty:
            ctx.append(estilo.estado_html("Nadie más del equipo sale", "ok"))
        else:
            nombres = ", ".join(estilo.nombre_pila(n) for n in cruces["nombre"].unique()[:3])
            ctx.append(estilo.estado_html(f"Se cruza con {nombres}", "no"))
        libres = saldo_num(r.cod)
        if libres is not None:
            ctx.append(estilo.estado_html(f"Le quedarían {libres} {'día' if libres == 1 else 'días'}",
                                          "espera" if libres <= 5 else "info"))
        if r.estado == "PENDIENTE_GERENCIA":
            ctx.append(estilo.estado_html("La aprueba gerencia", "espera"))
        if faltan < 0:
            ctx.append(estilo.estado_html("La fecha de inicio ya pasó", "no"))
        clave = f"{'pedu' if faltan <= 7 else 'ped'}_{prefijo}{r.id}"
        with st.container(border=True, key=clave):
            st.markdown(
                f'<div class="a-pedido"><span class="h-iniciales">{html.escape(estilo.iniciales(r.nombre))}</span>'
                f'<div class="a-quien"><h4>{html.escape(bonito(r.nombre))}</h4>'
                f'<p>Vacaciones {rango(r.inicio, r.fin)}, {int(r.dias)} '
                f'{"día" if r.dias == 1 else "días"}. {cuando}.</p><div class="a-ctx">{"".join(ctx)}</div></div></div>',
                unsafe_allow_html=True)
            _, b1, b2 = st.columns([2.2, 1, 1])
            b1.button("Rechazar", key=f"{prefijo}re{r.id}", use_container_width=True,
                      on_click=resolver, args=(r.id, "RECHAZADA", r.nombre))
            b2.button("Aprobar", key=f"{prefijo}ap{r.id}", type="primary", use_container_width=True,
                      on_click=resolver, args=(r.id, "APROBADA", r.nombre))


def resolver_tele(ids, nuevo_estado, nombre):
    with db.motor().begin() as con:
        con.execute(
            db.text("UPDATE teletrabajo_solicitudes SET estado = :e, resuelto_por = :q WHERE id = :id"),
            [{"e": nuevo_estado, "q": quien, "id": int(i)} for i in ids],
        )
    palabra = "aprobado" if nuevo_estado == "APROBADO" else "rechazado"
    st.session_state["ap_aviso"] = f"Teletrabajo {palabra}: {nombre}, {len(ids)} {'día' if len(ids) == 1 else 'días'}."


MESES = ["enero", "febrero", "marzo", "abril", "mayo", "junio", "julio",
         "agosto", "septiembre", "octubre", "noviembre", "diciembre"]


def semanas_del_mes(anio, mes):
    """[(lunes, [días hábiles del mes])] de cada semana con días hábiles del mes."""
    semanas = []
    for semana in calendar.Calendar(firstweekday=0).monthdatescalendar(anio, mes):
        habiles = [d for d in semana if d.month == mes and d.weekday() < 5]
        if habiles:
            semanas.append((semana[0], habiles))
    return semanas


def tarjetas_tele(df, prefijo):
    """Tabla por mes con la programación en espera: una fila por persona y una columna por semana.
    Las filas siguen el orden en que se registró la programación."""
    if df.empty:
        st.info("No hay teletrabajo por aprobar.")
        return
    df = df.assign(mes=df["fecha"].str[:7], creado=df["creado"].fillna(""))
    for mes_iso, grupo in df.groupby("mes"):
        anio, mes = int(mes_iso[:4]), int(mes_iso[5:])
        semanas = semanas_del_mes(anio, mes)
        cols_sem = [f"Semana {i + 1}\n{h[0]:%d/%m} al {h[-1]:%d/%m}" for i, (_, h) in enumerate(semanas)]
        orden = grupo.groupby("cod").agg(creado=("creado", "min"), nombre=("nombre", "first")) \
                     .sort_values(["creado", "nombre"]).index
        filas = []
        for cod in orden:
            g = grupo[grupo["cod"] == cod]
            dias = [date.fromisoformat(f) for f in g["fecha"]]
            fila = {"Elegir": True, "Colaborador": bonito(g.iloc[0]["nombre"]), "Puesto": bonito(g.iloc[0]["puesto"])}
            for col, (lun, _) in zip(cols_sem, semanas):
                fila[col] = ", ".join(f"{DIAS_SEMANA[d.weekday()]} {d:%d/%m}"
                                      for d in sorted(dias) if lun <= d < lun + timedelta(days=7))
            fila["cod"] = cod
            filas.append(fila)

        st.markdown(f"**{MESES[mes - 1].capitalize()} {anio}**: {len(filas)} "
                    f"{'persona' if len(filas) == 1 else 'personas'}, {len(grupo)} "
                    f"{'día' if len(grupo) == 1 else 'días'} en espera")
        editado = st.data_editor(
            pd.DataFrame(filas),
            key=f"{prefijo}tt{mes_iso}",
            hide_index=True,
            width="stretch",
            column_order=["Elegir", "Colaborador", "Puesto", *cols_sem],
            disabled=["Colaborador", "Puesto", *cols_sem],
            column_config={"Elegir": st.column_config.CheckboxColumn("Elegir", width="small",
                                                                     help="Desmarca a quien no quieras incluir")},
        )
        elegidos = editado[editado["Elegir"]]
        ids = grupo[grupo["cod"].isin(elegidos["cod"])]["id"].tolist()
        b1, b2, _ = st.columns([2, 2, 4])
        texto = f"{len(elegidos)} {'persona' if len(elegidos) == 1 else 'personas'}"
        if b1.button(f"Aprobar ({texto})", key=f"{prefijo}tap{mes_iso}", type="primary", disabled=not ids):
            resolver_tele(ids, "APROBADO", texto)
            st.rerun()
        if b2.button(f"Rechazar ({texto})", key=f"{prefijo}tre{mes_iso}", disabled=not ids):
            resolver_tele(ids, "RECHAZADO", texto)
            st.rerun()


def saldo(p):
    try:
        ingreso = date.fromisoformat(str(p.f_ingreso))
    except ValueError:
        return "Sin fecha de ingreso"
    c = inicio_ciclo(ingreso, hoy)
    if c is None:
        return f"Desde el {primer_aniversario(ingreso):%d/%m/%Y}"
    u = usados_en_ciclo(activas, p.cod, c, sumar_meses(c, MESES_CICLO))
    libres = 2 * DIAS_BLOQUE - u[1] - u[2]
    return f"{libres} {'día' if libres == 1 else 'días'}"


def proximas(cod):
    futuras = activas[(activas["cod"] == cod) & (activas["fin"] >= hoy.isoformat())]
    if futuras.empty:
        return ""
    r = futuras.iloc[0]
    return f"{fecha(r['inicio'], False)} al {fecha(r['fin'], False)} ({ESTADOS[r['estado']].lower()})"


def proximo_corto(cod):
    futuras = activas[(activas["cod"] == cod) & (activas["fin"] >= hoy.isoformat())].sort_values("inicio")
    if futuras.empty:
        return "Sin programar"
    r = futuras.iloc[0]
    d = date.fromisoformat(r["inicio"])
    estado = {"APROBADA": "aprobada", "PENDIENTE": "en espera", "PENDIENTE_GERENCIA": "en espera de gerencia"}.get(r["estado"], "")
    return f"{d.day} {estilo.MESES[d.month - 1][:3]}<small>{estado}</small>"


def tabla_personal(equipo):
    filas = []
    for p in equipo.itertuples():
        libres = saldo_num(p.cod)
        if libres is None:
            try:
                desde = f"Desde {primer_aniversario(date.fromisoformat(str(p.f_ingreso))):%m/%Y}"
            except ValueError:
                desde = "Sin fecha de ingreso"
            saldo_html = f'<div class="a-barra"><i style="width:0"></i></div><small>{desde}</small>'
        else:
            saldo_html = (f'<div class="a-barra"><i style="width:{max(0, min(100, libres * 100 // 30))}%"></i></div>'
                          f"<small>{libres} de 30</small>")
        filas.append(f"<tr><td>{html.escape(bonito(p.nombre))}<small>{html.escape(bonito(p.puesto))}</small></td>"
                     f"<td>{saldo_html}</td><td>{proximo_corto(p.cod)}</td></tr>")
    st.markdown('<div class="h-panel"><h3>Tu equipo</h3><table class="a-equipo"><thead><tr><th>Colaborador</th>'
                '<th style="width:30%">Días por programar</th><th style="width:22%">Próximo</th></tr></thead><tbody>'
                + "".join(filas) + "</tbody></table></div>", unsafe_allow_html=True)


def seccion_area(area):
    equipo = personal[(personal["area"] == area) & (personal["cod"] != mi_cod)]
    pend = sols[(sols["area"] == area) & (sols["estado"] == "PENDIENTE")
                & (sols["cod"] != mi_cod) & ~sols["cod"].isin(cods_jefes)]
    tpend = tele_pend[(tele_pend["area"] == area) & (tele_pend["cod"] != mi_cod)
                      & ~tele_pend["cod"].isin(cods_jefes)]
    total = len(pend) + tpend["cod"].nunique()
    st.markdown(f'<div class="h-sub">{bonito(area)}: '
                + (f'{total} {"solicitud de tu equipo espera" if total == 1 else "solicitudes de tu equipo esperan"} tu respuesta.'
                   if total else "no hay solicitudes pendientes.") + "</div>", unsafe_allow_html=True)
    izq, der = st.columns([1.2, 1], gap="medium")
    with izq:
        st.markdown('<div class="a-cab"><b>Vacaciones por aprobar</b><span>La más próxima, primero</span></div>',
                    unsafe_allow_html=True)
        tarjetas(pend, f"{area[:6]}_")
    with der:
        if equipo.empty:
            st.info("No hay personal cargado en esta área.")
        else:
            tabla_personal(equipo)
    st.markdown('<div class="a-cab" style="margin-top:18px"><b>Teletrabajo por aprobar</b>'
                '<span>Una fila por persona, en el orden en que se programó</span></div>', unsafe_allow_html=True)
    tarjetas_tele(tpend, f"{area[:6]}_")


# ---------- Jefe (o admin que también es jefe de un área) ----------
if rol == "jefe" and not mis_areas:
    st.warning("No figuras como jefe de ningún área. Pide al administrador que revise la lista de jefes por área.")
    st.stop()

for area in mis_areas:
    seccion_area(area)

if rol != "admin":
    st.stop()

# ---------- Administrador ----------
if mis_areas:
    st.divider()
total_admin = int(sols["estado"].str.startswith("PENDIENTE").sum()) + tele_pend["cod"].nunique()
st.markdown(f'<div class="h-sub">Gerencia, jefaturas y áreas sin jefe. En total hay {total_admin} '
            f'{"solicitud" if total_admin == 1 else "solicitudes"} en espera.</div>', unsafe_allow_html=True)

no_encontrados = [f"{bonito(a)} ({ap})" for a, ap in APROBADORES.items() if a not in jefes]
if no_encontrados:
    st.warning("No encontré en el personal al jefe de: " + ", ".join(no_encontrados) + ". "
               "Sus solicitudes te llegan a ti.")
areas_sin_jefe = sorted(set(personal["area"].dropna()) - set(jefes))
if areas_sin_jefe:
    st.caption("Áreas sin jefe asignado (las apruebas tú): " + ", ".join(bonito(a) for a in areas_sin_jefe) + ".")

st.markdown('<div class="a-cab" style="margin-top:14px"><b>Requieren gerencia</b></div>', unsafe_allow_html=True)
tarjetas(sols[sols["estado"] == "PENDIENTE_GERENCIA"], "g_")

st.markdown('<div class="a-cab" style="margin-top:14px"><b>Jefaturas y áreas sin jefe</b></div>', unsafe_allow_html=True)
pend = sols[sols["estado"] == "PENDIENTE"]
tarjetas(pend[(~pend["area"].isin(list(jefes))) | (pend["cod"].isin(cods_jefes))], "a_")

directo = (~tele_pend["area"].isin(list(jefes))) | (tele_pend["cod"].isin(cods_jefes))
if directo.any():
    st.markdown('<div class="a-cab" style="margin-top:14px"><b>Teletrabajo: jefaturas y áreas sin jefe</b></div>', unsafe_allow_html=True)
    tarjetas_tele(tele_pend[directo], "a_")

with st.expander("Pendientes que corresponden a cada jefe"):
    otras = pend[pend["area"].isin(list(jefes)) & ~pend["cod"].isin(cods_jefes) & ~pend["area"].isin(mis_areas)]
    st.markdown("**Vacaciones**")
    tarjetas(otras, "o_")
    st.markdown("**Teletrabajo**")
    tarjetas_tele(tele_pend[~directo & ~tele_pend["area"].isin(mis_areas)], "o_")

with st.expander("Todas las solicitudes"):
    vista = sols.sort_values("inicio", ascending=False)
    st.dataframe(
        vista.assign(
            Colaborador=vista["nombre"].map(bonito), Área=vista["area"].map(bonito),
            Desde=vista["inicio"].map(fecha), Hasta=vista["fin"].map(fecha),
            Estado=vista["estado"].map(ESTADOS), Respondió=vista["resuelto_por"].fillna("").map(
                lambda x: bonito(x) if x else ""),
        )[["Colaborador", "Área", "Desde", "Hasta", "dias", "Estado", "Respondió"]].rename(columns={"dias": "Días"}),
        width="stretch", hide_index=True,
    )


# ---------- Enviar vacaciones por correo ----------
def excel_vacaciones(df):
    import io
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    wb = Workbook()
    ws = wb.active
    ws.title = "Vacaciones"
    cols = ["Colaborador", "Puesto", "Área", "Desde", "Hasta", "Días", "Estado", "Respondió"]
    for i, (t, ancho) in enumerate(zip(cols, [34, 26, 34, 12, 12, 7, 20, 30]), start=1):
        c = ws.cell(row=1, column=i, value=t)
        c.font, c.fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", start_color="0B4F9C")
        c.alignment = Alignment(horizontal="center")
        ws.column_dimensions[c.column_letter].width = ancho
    for f, r in enumerate(df.itertuples(), start=2):
        valores = [bonito(r.nombre), bonito(r.puesto), bonito(r.area), date.fromisoformat(r.inicio),
                   date.fromisoformat(r.fin), int(r.dias), ESTADOS.get(r.estado, r.estado),
                   bonito(r.resuelto_por) if isinstance(r.resuelto_por, str) and r.resuelto_por else ""]
        for col, v in enumerate(valores, start=1):
            c = ws.cell(row=f, column=col, value=v)
            if col in (4, 5):
                c.number_format = "dd/mm/yyyy"
    ws.freeze_panes = "A2"
    if len(df):
        ws.auto_filter.ref = f"A1:H{len(df) + 1}"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def html_vacaciones(df, titulo):
    celda = "padding:6px 10px;border-bottom:1px solid #D5E2F3;text-align:left"
    resumen = df["estado"].map(ESTADOS).value_counts()
    chips = " &nbsp;·&nbsp; ".join(f"<b>{n}</b> {html.escape(e.lower())}" for e, n in resumen.items())

    def td(v, derecha=False):
        return f"<td style='{celda}{';text-align:right' if derecha else ''}'>{html.escape(str(v))}</td>"

    filas = "".join(
        "<tr>" + td(bonito(r.nombre)) + td(bonito(r.area)) + td(fecha(r.inicio)) + td(fecha(r.fin))
        + td(int(r.dias), derecha=True) + td(ESTADOS.get(r.estado, r.estado)) + "</tr>"
        for r in df.head(300).itertuples())
    cabecera = "".join(f"<th style='{celda}'>{t}</th>" for t in ("Colaborador", "Área", "Desde", "Hasta", "Días", "Estado"))
    extra = f"<p>Se muestran 300 de {len(df)}; el resto está en el Excel adjunto.</p>" if len(df) > 300 else ""
    return f"""<div style="font-family:Segoe UI,Arial,sans-serif;color:#14294B">
<h2 style="color:#0B4F9C;margin:0 0 4px">{html.escape(titulo)}</h2>
<p style="margin:0 0 12px;color:#4A5D7E">{len(df)} solicitudes: {chips}</p>
<table style="border-collapse:collapse;font-size:13px">
<tr style="background:#0B4F9C;color:#FFFFFF">{cabecera}</tr>
{filas}
</table>{extra}
<p style="color:#6A7C99;font-size:12px">Enviado desde HERMES el {hoy:%d/%m/%Y}. El detalle completo va en el Excel adjunto.</p>
</div>"""


with st.expander("Enviar vacaciones por correo"):
    c1, c2 = st.columns([2, 1])
    alcance = c1.radio("¿Qué vacaciones?", ["Próximas (de hoy en adelante)", "Todas"], horizontal=True,
                       key="mail_alcance")
    con_rechazadas = c2.checkbox("Incluir rechazadas", value=False, key="mail_rech")
    envio = sols.copy()
    if alcance.startswith("Próximas"):
        envio = envio[envio["fin"] >= hoy.isoformat()]
    if not con_rechazadas:
        envio = envio[envio["estado"] != "RECHAZADA"]
    envio = envio.sort_values(["inicio", "nombre"])
    st.caption(f"Se enviarán {len(envio)} solicitudes, con el Excel adjunto.")

    destino = st.text_input("Enviar a (puedes poner varios, separados por coma)",
                            value=correo.destino_por_defecto(), key="mail_destino")
    archivo_excel = excel_vacaciones(envio)
    nombre_excel = f"vacaciones_{hoy:%Y-%m-%d}.xlsx"
    b1, b2 = st.columns(2)
    if not correo.configurado():
        st.info("Para enviar correos falta configurar la cuenta de Outlook en los secrets "
                "(SMTP_USUARIO y SMTP_CLAVE). Mientras tanto, puedes descargar el Excel.")
    if b1.button("Enviar correo", type="primary", disabled=envio.empty or not correo.configurado()):
        with st.spinner("Enviando..."):
            titulo = ("Vacaciones programadas" if alcance.startswith("Próximas") else "Todas las vacaciones") \
                     + f" al {hoy:%d/%m/%Y}"
            error = correo.enviar(destino, f"HERMES · {titulo}", html_vacaciones(envio, titulo),
                                  [(nombre_excel, archivo_excel,
                                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")])
        if error:
            st.error(error)
        else:
            st.success(f"Correo enviado a {destino}.")
    b2.download_button("Descargar el Excel", data=archivo_excel, file_name=nombre_excel,
                       mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                       disabled=envio.empty, key="mail_excel")
