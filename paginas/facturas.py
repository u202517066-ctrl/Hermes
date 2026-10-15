import html
import io
from datetime import date, timedelta

import pandas as pd
import streamlit as st
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from sqlalchemy import text

import db
import estilo
from lector_facturas import VERSION, clave_api, leer_archivo

st.markdown(
    """
<style>
.fa-banda {
  background: #0B4F9C; color: #FFFFFF; border-radius: 14px;
  padding: 1.2rem 1.6rem; display: flex; justify-content: space-between;
  align-items: flex-end; gap: 1rem; flex-wrap: wrap; margin-bottom: 0.8rem;
}
.fa-banda h2 { color: #FFFFFF; margin: 0; padding: 0; font-size: 1.6rem; font-weight: 600; }
.fa-banda p { margin: 0.25rem 0 0; opacity: 0.85; }
.fa-num { text-align: right; line-height: 1.1; }
.fa-num b { font-size: 2.4rem; display: block; }
.fa-num span { opacity: 0.85; font-size: 0.95rem; }
/* Sin botones − / +: los montos leídos de la factura solo se corrigen escribiendo */
[data-testid="stNumberInputStepUp"], [data-testid="stNumberInputStepDown"] { display: none; }
</style>
""",
    unsafe_allow_html=True,
)

DIAS_ALERTA = 7  # avisar las que vencen en los próximos 7 días
SIMBOLO = {"PEN": "S/", "USD": "US$"}
# Flujo: Por aprobar -> Aprobada (u Observada) -> Pagada
ESTADOS = {"PENDIENTE": "Por aprobar", "APROBADA": "Aprobada", "OBSERVADA": "Observada",
           "PAGADA": "Pagada", "ANULADA": "Anulada"}
NO_PAGADAS = ("PENDIENTE", "APROBADA", "OBSERVADA")

st.title("Facturas de proveedores")
sesion = st.session_state.get("sesion", {})
roles = set(sesion.get("roles") or [sesion.get("rol")])
# Vista de administrador si es admin; si no, vista de aprobador (aunque además sea jefe o colaborador)
rol = "admin" if "admin" in roles else "aprobador_facturas" if "aprobador_facturas" in roles else None
puede_aprobar = "aprobador_facturas" in roles
if rol not in ("admin", "aprobador_facturas"):
    st.error("Solo el administrador y el aprobador de facturas pueden ver esta página.")
    st.stop()

hoy = db.hoy()
st.session_state.setdefault("fa_ronda", 0)  # cambia tras cada acción para limpiar las casillas marcadas


def listo(mensaje):
    st.session_state["fa_aviso"] = mensaje
    st.session_state["fa_ronda"] += 1
    st.rerun()


if "fa_aviso" in st.session_state:
    st.success(st.session_state.pop("fa_aviso"))


# ---------- Utilidades ----------
def dinero(monto, moneda):
    return f"{SIMBOLO.get(moneda, moneda or '')} {monto:,.2f}".strip()


def totales_por_moneda(df):
    if df.empty:
        return "S/ 0.00"
    return " y ".join(dinero(v, m) for m, v in df.groupby("moneda")["total"].sum().items())


def a_fecha(iso):
    """Convierte 'AAAA-MM-DD' a fecha; vacío o sin dato -> None."""
    return date.fromisoformat(iso) if isinstance(iso, str) and iso else None


def fecha_corta(iso):
    return date.fromisoformat(iso).strftime("%d/%m/%Y") if isinstance(iso, str) and iso else ""


def cargar():
    df = db.consultar("SELECT * FROM facturas_proveedores ORDER BY fecha_vencimiento, proveedor")
    for c in ("subtotal", "igv", "total"):
        df[c] = pd.to_numeric(df[c], errors="coerce").fillna(0.0)
    df["moneda"] = df["moneda"].fillna("PEN")
    for c in ("aprobado_por", "fecha_aprobacion", "motivo_observacion"):
        if c not in df.columns:
            df[c] = None
    df["vencida"] = df["estado"].isin(NO_PAGADAS) & (df["fecha_vencimiento"].fillna("9999") < hoy.isoformat())
    df["situacion"] = df["estado"].map(ESTADOS).fillna(df["estado"])
    df["dias"] = [(date.fromisoformat(v) - hoy).days if isinstance(v, str) and v else None
                  for v in df["fecha_vencimiento"]]
    df["alerta"] = [alerta(e, d) for e, d in zip(df["estado"], df["dias"])]
    return df


def alerta(estado, dias, iconos=True):
    """Texto de alerta según los días que faltan para el vencimiento (solo facturas no pagadas)."""
    if estado not in NO_PAGADAS or dias is None or pd.isna(dias):
        return ""
    dias = int(dias)
    rojo, ambar = ("🔴 ", "🟡 ") if iconos else ("", "")
    if dias < 0:
        return f"{rojo}Vencida hace {-dias} {'día' if dias == -1 else 'días'}"
    if dias == 0:
        return f"{rojo}Vence hoy"
    if dias <= DIAS_ALERTA:
        return f"{ambar}Vence en {dias} {'día' if dias == 1 else 'días'}"
    return ""


def guardar(filas):
    """Guarda las facturas nuevas con su detalle y omite las repetidas (mismo RUC y comprobante).
    Cada fila puede traer 'items': [{descripcion, cantidad, valor_unitario, importe}]. Devuelve (guardadas, repetidas)."""
    existentes = db.consultar("SELECT ruc, comprobante FROM facturas_proveedores")
    claves = {(str(r or "").strip(), str(c or "").strip().upper())
              for r, c in zip(existentes["ruc"], existentes["comprobante"])}
    guardadas, repetidas = 0, []
    with db.motor().begin() as con:
        for f in filas:
            clave = (str(f["ruc"] or "").strip(), str(f["comprobante"] or "").strip().upper())
            if clave in claves:
                repetidas.append(f["comprobante"])
                continue
            claves.add(clave)
            datos = {k: v for k, v in f.items() if k not in ("items", "documento")}
            factura_id = con.execute(text(
                "INSERT INTO facturas_proveedores (ruc, proveedor, comprobante, fecha_emision, fecha_vencimiento, "
                "moneda, subtotal, igv, total, estado, observacion, archivo, creado) VALUES (:ruc, :proveedor, "
                ":comprobante, :fecha_emision, :fecha_vencimiento, :moneda, :subtotal, :igv, :total, 'PENDIENTE', "
                ":observacion, :archivo, :creado) RETURNING id"), {**datos, "creado": db.ahora()}).scalar()
            if f.get("documento"):
                guardar_documento(con, factura_id, *f["documento"])
            items = [it for it in f.get("items", []) if it.get("descripcion") or it.get("importe")]
            if items:
                con.execute(text(
                    "INSERT INTO facturas_items (factura_id, n, descripcion, cantidad, valor_unitario, importe) "
                    "VALUES (:factura_id, :n, :descripcion, :cantidad, :valor_unitario, :importe)"),
                    [{**it, "factura_id": factura_id, "n": i} for i, it in enumerate(items, start=1)])
            guardadas += 1
    return guardadas, repetidas


# ---------- Documento original de la factura ----------
TIPOS = {".pdf": "application/pdf", ".xml": "application/xml", ".zip": "application/zip",
         ".png": "image/png", ".jpg": "image/jpeg", ".jpeg": "image/jpeg", ".webp": "image/webp"}


def tipo_de(nombre):
    return TIPOS.get("." + nombre.lower().rsplit(".", 1)[-1], "application/octet-stream")


def guardar_documento(con, factura_id, nombre, contenido):
    con.execute(text("DELETE FROM facturas_archivos WHERE factura_id = :f"), {"f": factura_id})
    con.execute(text("INSERT INTO facturas_archivos (factura_id, nombre, tipo, contenido) VALUES (:f, :n, :t, :c)"),
                {"f": factura_id, "n": nombre, "t": tipo_de(nombre), "c": contenido})


def documento_de(factura_id):
    fila = db.uno("SELECT nombre, contenido FROM facturas_archivos WHERE factura_id = :f", {"f": int(factura_id)})
    return (fila.nombre, bytes(fila.contenido)) if fila else None


@st.cache_data(show_spinner=False, max_entries=60)
def paginas_pdf(contenido, maximo=4):
    import pymupdf
    imagenes = []
    with pymupdf.open(stream=contenido, filetype="pdf") as doc:
        for pagina in list(doc)[:maximo]:
            imagenes.append(pagina.get_pixmap(dpi=110).tobytes("png"))
        total = len(doc)
    return imagenes, total


def mostrar_documento(nombre, contenido, clave):
    """Muestra el documento original en pantalla, con un botón para descargarlo."""
    import io as _io
    import zipfile
    ext = nombre.lower().rsplit(".", 1)[-1]
    st.download_button("Descargar documento", data=contenido, file_name=nombre, mime=tipo_de(nombre),
                       key=f"desc_{clave}")
    try:
        if ext == "zip":
            with zipfile.ZipFile(_io.BytesIO(contenido)) as z:
                internos = [n for n in z.namelist() if n.lower().endswith((".pdf", ".xml"))]
                pdfs = [n for n in internos if n.lower().endswith(".pdf")]
                elegido = pdfs[0] if pdfs else (internos[0] if internos else None)
                if not elegido:
                    st.caption("El ZIP no trae PDF ni XML para mostrar.")
                    return
                st.caption(f"Mostrando {elegido}, dentro del ZIP.")
                nombre, contenido, ext = elegido, z.read(elegido), elegido.lower().rsplit(".", 1)[-1]
        if ext == "pdf":
            imagenes, total = paginas_pdf(contenido)
            for img in imagenes:
                st.image(img)
            if total > len(imagenes):
                st.caption(f"Se muestran {len(imagenes)} de {total} páginas. Descarga el documento para verlo completo.")
        elif ext in ("png", "jpg", "jpeg", "webp"):
            st.image(contenido)
        elif ext == "xml":
            st.code(contenido.decode("utf-8", errors="replace")[:30000], language="xml")
        else:
            st.caption("Este tipo de archivo no se puede mostrar; descárgalo para verlo.")
    except Exception:
        st.caption("No se pudo mostrar el documento; descárgalo para verlo.")


def ver_documento_guardado(factura_id, clave):
    doc = documento_de(factura_id)
    if doc:
        mostrar_documento(doc[0], doc[1], clave)
    else:
        st.caption("Esta factura no tiene documento guardado (se registró antes de esta mejora o a mano). "
                   "Puedes adjuntarlo al editarla.")


def reemplazar_items(con, factura_id, items):
    con.execute(text("DELETE FROM facturas_items WHERE factura_id = :f"), {"f": factura_id})
    items = [it for it in items if it.get("descripcion") or it.get("importe")]
    if items:
        con.execute(text("INSERT INTO facturas_items (factura_id, n, descripcion, cantidad, valor_unitario, importe) "
                         "VALUES (:f, :n, :descripcion, :cantidad, :valor_unitario, :importe)"),
                    [{**it, "f": factura_id, "n": i} for i, it in enumerate(items, start=1)])


def releer_facturas(ids):
    """Vuelve a leer el documento guardado de cada factura con el lector actual y actualiza sus datos y servicios.
    Si estaba aprobada u observada, vuelve a aprobación porque sus datos pueden cambiar.
    Devuelve (releídas, sin_documento)."""
    releidas, sin_doc = 0, 0
    actuales = cargar().set_index("id")
    for fid in ids:
        doc = documento_de(fid)
        if not doc:
            sin_doc += 1
            continue
        datos, _ = leer_archivo(doc[0], doc[1])
        f = actuales.loc[fid]
        nuevo = {c: (datos[c] if datos[c] not in (None, "") else f[c]) for c in
                 ("ruc", "proveedor", "comprobante", "fecha_emision", "fecha_vencimiento", "moneda",
                  "subtotal", "igv", "total")}
        if not nuevo["fecha_vencimiento"]:
            nuevo["fecha_vencimiento"] = nuevo["fecha_emision"]
        for c in ("subtotal", "igv", "total"):
            nuevo[c] = None if nuevo[c] is None or pd.isna(nuevo[c]) else float(nuevo[c])
        vuelve = f["estado"] in ("APROBADA", "OBSERVADA")
        with db.motor().begin() as con:
            con.execute(text(
                "UPDATE facturas_proveedores SET ruc = :ruc, proveedor = :proveedor, comprobante = :comprobante, "
                "fecha_emision = :fecha_emision, fecha_vencimiento = :fecha_vencimiento, moneda = :moneda, "
                "subtotal = :subtotal, igv = :igv, total = :total" + (
                    ", estado = 'PENDIENTE', aprobado_por = NULL, fecha_aprobacion = NULL, motivo_observacion = NULL"
                    if vuelve else "") + " WHERE id = :id"), {**nuevo, "id": int(fid)})
            if datos["items"]:
                reemplazar_items(con, int(fid), datos["items"])
        releidas += 1
    return releidas, sin_doc


def cargar_items(ids=None):
    df = db.consultar("SELECT * FROM facturas_items ORDER BY factura_id, n")
    for c in ("cantidad", "valor_unitario", "importe"):
        df[c] = pd.to_numeric(df[c], errors="coerce")
    return df if ids is None else df[df["factura_id"].isin(ids)]


@st.cache_data(show_spinner=False, max_entries=300)
def leer(nombre, contenido, version):
    """Se guarda en memoria para no releer cada vez que se edita una celda.
    'version' hace que, al actualizar el lector, las facturas se vuelvan a leer."""
    return leer_archivo(nombre, contenido)


def limpiar_items(df):
    items = []
    for r in df.to_dict("records"):
        v = {k: (None if pd.isna(x) else x) for k, x in r.items()}
        if v["Descripción"] or v["Importe"]:
            items.append({"descripcion": v["Descripción"], "cantidad": v["Cantidad"],
                          "valor_unitario": v["Valor unitario"], "importe": v["Importe"]})
    return items


COLS_ITEMS = {
    "Descripción": st.column_config.TextColumn("Descripción", width="large"),
    "Cantidad": st.column_config.NumberColumn("Cantidad", format="%.2f"),
    "Valor unitario": st.column_config.NumberColumn("Valor unitario", format="%.2f"),
    "Importe": st.column_config.NumberColumn("Importe", format="%.2f"),
}


def tabla_items(items):
    return pd.DataFrame(items or [{"descripcion": None, "cantidad": None, "valor_unitario": None, "importe": None}]) \
        .rename(columns={"descripcion": "Descripción", "cantidad": "Cantidad",
                         "valor_unitario": "Valor unitario", "importe": "Importe"})[list(COLS_ITEMS)]


def excel(df):
    """Excel con una fila por servicio: los datos de la factura se repiten en cada fila de sus servicios.
    Subtotal, IGV y Total van solo en la primera fila de cada factura, para que las sumas no se dupliquen."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Facturas"
    azul, celeste, borde = "0B4F9C", "D6E4F5", Side(style="thin", color="D5E2F3")
    columnas = ["Proveedor", "RUC", "Comprobante", "Fecha emisión", "Fecha vencimiento", "Moneda", "Estado",
                "N.°", "Descripción del servicio", "Cantidad", "Valor unitario", "Importe",
                "Subtotal factura", "IGV factura", "Total factura",
                "Aprobado por", "Fecha aprobación", "Fecha de pago", "Observación", "Alerta"]
    anchos = [32, 13, 16, 13, 13, 8, 18, 6, 50, 10, 13, 13, 14, 12, 14, 22, 13, 13, 30, 22]
    servicio = range(8, 13)  # columnas del servicio, en celeste
    fechas, montos = (4, 5, 17, 18), (10, 11, 12, 13, 14, 15)

    ws["A1"] = "FACTURAS DE PROVEEDORES"
    ws["A1"].font = Font(bold=True, size=14, color=azul)
    ws["A2"] = f"Generado el {hoy:%d/%m/%Y}. Una fila por servicio; los montos de la factura van en su primera fila."
    ws["A2"].font = Font(italic=True, color="6A7C99")
    for i, (titulo, ancho) in enumerate(zip(columnas, anchos), start=1):
        c = ws.cell(row=4, column=i, value=titulo)
        if i in servicio:
            c.font, c.fill = Font(bold=True, color=azul), PatternFill("solid", start_color=celeste)
        else:
            c.font, c.fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", start_color=azul)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = ancho
    ws.row_dimensions[4].height = 30

    items = cargar_items(df["id"].tolist())
    por_factura = {fid: g.sort_values("n") for fid, g in items.groupby("factura_id")}
    fila = 5
    for r in df.itertuples():
        cabecera = [r.proveedor, r.ruc, r.comprobante, a_fecha(r.fecha_emision), a_fecha(r.fecha_vencimiento),
                    r.moneda, r.situacion]
        cierre = [r.aprobado_por, a_fecha(r.fecha_aprobacion), a_fecha(r.fecha_pago),
                  "; ".join(x for x in (r.observacion, f"Observada: {r.motivo_observacion}"
                                        if r.motivo_observacion else None) if x),
                  alerta(r.estado, r.dias, iconos=False)]
        servicios = por_factura.get(r.id)
        lineas = list(servicios.itertuples()) if servicios is not None else [None]
        for k, it in enumerate(lineas):
            detalle = [None] * 5 if it is None else [
                it.n, it.descripcion,
                None if pd.isna(it.cantidad) else float(it.cantidad),
                None if pd.isna(it.valor_unitario) else float(it.valor_unitario),
                None if pd.isna(it.importe) else float(it.importe)]
            montos_factura = [float(r.subtotal), float(r.igv), float(r.total)] if k == 0 else [None] * 3
            for col, v in enumerate(cabecera + detalle + montos_factura + cierre, start=1):
                c = ws.cell(row=fila, column=col, value=v)
                c.border = Border(bottom=borde)
                c.alignment = Alignment(vertical="top", wrap_text=(col == 9))
                if col in fechas:
                    c.number_format = "dd/mm/yyyy"
                elif col in montos:
                    c.number_format = "#,##0.00"
            fila += 1

    ultima = fila - 1
    if ultima >= 5:
        fila += 1
        for moneda in sorted(df["moneda"].unique()):
            ws.cell(row=fila, column=14, value=f"Total {moneda}").font = Font(bold=True)
            c = ws.cell(row=fila, column=15, value=f'=SUMIF(F5:F{ultima},"{moneda}",O5:O{ultima})')
            c.font, c.number_format = Font(bold=True), "#,##0.00"
            fila += 1
        ws.auto_filter.ref = f"A4:{get_column_letter(len(columnas))}{ultima}"
    ws.freeze_panes = "D5"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---------- Datos ----------
facturas = cargar()
por_aprobar = facturas[facturas["estado"] == "PENDIENTE"]
aprobadas = facturas[facturas["estado"] == "APROBADA"]
observadas = facturas[facturas["estado"] == "OBSERVADA"]


def nombre_actual():
    u = db.uno("SELECT nombre, apellido FROM usuarios WHERE usuario = :u", {"u": sesion.get("usuario")})
    nombre = f"{u.nombre or ''} {u.apellido or ''}".strip() if u else ""
    return nombre or sesion.get("usuario")


st.markdown(
    """
<style>
.f-deuda { background: var(--h-superficie); border: 1px solid var(--h-linea); border-radius: 16px; padding: 18px 22px; margin: 6px 0 16px; }
.f-deuda-cab { display: flex; justify-content: space-between; align-items: baseline; gap: 12px; flex-wrap: wrap; }
.f-deuda-cab b { font-stretch: 75%; font-size: 2.5rem; font-weight: 800; line-height: 1; color: var(--h-tinta); }
.f-deuda-cab span { color: var(--h-suave); }
.f-barra { display: flex; height: 16px; border-radius: 99px; overflow: hidden; margin: 14px 0 10px; gap: 2px; background: var(--h-fondo); }
.f-barra i { display: block; height: 100%; }
.f-ley { display: flex; gap: 26px; flex-wrap: wrap; font-size: 0.9rem; color: var(--h-suave); }
.f-ley div::before { content: ""; display: inline-block; width: 10px; height: 10px; border-radius: 3px; margin-right: 7px; }
.f-ley .v::before { background: var(--h-rojo); } .f-ley .p::before { background: var(--h-ambar); } .f-ley .a::before { background: var(--h-azul); }
.f-ley strong { display: block; font-size: 1.05rem; color: var(--h-tinta); }
.f-ficha h3 { margin: 0; font-size: 1.15rem; color: var(--h-tinta); }
.f-ficha .sub { color: var(--h-suave); margin: 2px 0 12px; font-size: 0.9rem; }
.f-serv { list-style: none; margin: 4px 0 14px; padding: 0; }
.f-serv li { display: flex; justify-content: space-between; gap: 12px; padding: 7px 0; border-bottom: 1px dashed var(--h-linea);
  font-size: 0.9rem; color: var(--h-tinta); }
.f-serv li span:last-child { font-variant-numeric: tabular-nums; font-weight: 600; white-space: nowrap; }
.f-serv li.total { border-bottom: 0; font-weight: 700; }
.f-pasos { list-style: none; margin: 0 0 12px; padding: 0; display: grid; grid-template-columns: repeat(3, 1fr); gap: 6px; }
.f-pasos li { font-size: 0.8rem; color: var(--h-suave); padding-top: 9px; border-top: 4px solid var(--h-linea); }
.f-pasos li.hecho { border-top-color: var(--h-verde); color: var(--h-tinta); }
.f-pasos li.ahora { border-top-color: var(--h-ambar); color: var(--h-tinta); }
.f-pasos li.mal { border-top-color: var(--h-rojo); color: var(--h-tinta); }
.f-pasos li b { display: block; font-size: 0.85rem; }
.f-vacio { color: var(--h-suave); text-align: center; padding: 40px 10px; }
.st-key-ficha { background: var(--h-superficie); }
[data-testid="stPills"] button[aria-checked="true"], [data-testid="stPills"] button[kind="pillsActive"] {
  background: var(--h-celeste) !important; border-color: var(--h-azul) !important; color: var(--h-azul) !important; }
</style>
""",
    unsafe_allow_html=True,
)

subt = (f"{len(por_aprobar)} por aprobar" + (f" y {len(aprobadas)} {'aprobada' if len(aprobadas) == 1 else 'aprobadas'} "
        "esperando pago" if rol == "admin" else "") + "."
        + (f" {len(observadas)} {'observada' if len(observadas) == 1 else 'observadas'} por corregir." if rol == "admin" and len(observadas) else ""))
st.markdown(f'<div class="h-sub">{html.escape(subt)}</div>', unsafe_allow_html=True)

if rol == "admin":
    tab_main, tab_cargar, tab_mano = st.tabs(["Facturas", "Cargar facturas", "Agregar a mano"])
else:
    tab_main = st.container()


def editar_factura(f):
    fid = int(f["id"])
    with st.container(border=True):
        st.markdown(f"**Editar {f['comprobante']}**")
        e1, e2, e3 = st.columns([3, 2, 2])
        n_prov = e1.text_input("Proveedor", f["proveedor"] or "", key=f"ed_p_{fid}")
        n_ruc = e2.text_input("RUC", f["ruc"] or "", key=f"ed_r_{fid}")
        n_comp = e3.text_input("Comprobante", f["comprobante"] or "", key=f"ed_c_{fid}")
        e4, e5, e6 = st.columns(3)
        n_emi = e4.date_input("Emisión", (a_fecha(f["fecha_emision"]) or hoy), format="DD/MM/YYYY", key=f"ed_e_{fid}")
        n_ven = e5.date_input("Vence", (a_fecha(f["fecha_vencimiento"]) or hoy), format="DD/MM/YYYY", key=f"ed_v_{fid}")
        n_mon = e6.selectbox("Moneda", ["PEN", "USD"], index=1 if f["moneda"] == "USD" else 0, key=f"ed_m_{fid}")
        e7, e8, e9 = st.columns(3)
        n_sub = e7.number_input("Subtotal", value=float(f["subtotal"]), min_value=0.0, format="%.2f", key=f"ed_s_{fid}")
        n_igv = e8.number_input("IGV", value=float(f["igv"]), min_value=0.0, format="%.2f", key=f"ed_i_{fid}")
        n_tot = e9.number_input("Total", value=float(f["total"]), min_value=0.0, format="%.2f", key=f"ed_t_{fid}")
        n_obs = st.text_input("Observación", f["observacion"] or "", key=f"ed_o_{fid}")
        st.markdown("Detalle de servicios")
        n_items = st.data_editor(tabla_items(cargar_items([fid]).to_dict("records")), key=f"ed_it_{fid}",
                                 num_rows="dynamic", hide_index=True, width="stretch", column_config=COLS_ITEMS)
        adjunto = st.file_uploader("Adjuntar o reemplazar el documento (opcional)",
                                   type=["pdf", "xml", "zip", "png", "jpg", "jpeg", "webp"], key=f"ed_doc_{fid}")
        g1, g2, _ = st.columns([1, 1, 2])
        if g2.button("Cancelar", key=f"ed_x_{fid}", use_container_width=True):
            st.session_state.pop("fa_editando", None)
            st.rerun()
        if g1.button("Guardar cambios", type="primary", key=f"ed_g_{fid}", use_container_width=True):
            if n_ven < n_emi:
                st.error("La fecha de vencimiento es anterior a la de emisión.")
            elif not n_prov.strip() or not n_comp.strip() or n_tot <= 0:
                st.error("Completa proveedor, comprobante y total.")
            else:
                vuelve = f["estado"] in ("APROBADA", "OBSERVADA")
                with db.motor().begin() as con:
                    if adjunto is not None:
                        guardar_documento(con, fid, adjunto.name, adjunto.getvalue())
                    reemplazar_items(con, fid, limpiar_items(n_items))
                db.ejecutar(
                    "UPDATE facturas_proveedores SET proveedor = :p, ruc = :r, comprobante = :c, fecha_emision = :e, "
                    "fecha_vencimiento = :v, moneda = :m, subtotal = :s, igv = :i, total = :t, observacion = :o" + (
                        ", estado = 'PENDIENTE', aprobado_por = NULL, fecha_aprobacion = NULL, motivo_observacion = NULL"
                        if vuelve else "") + " WHERE id = :id",
                    {"p": n_prov.strip(), "r": n_ruc.strip() or None, "c": n_comp.strip().upper(), "e": n_emi.isoformat(),
                     "v": n_ven.isoformat(), "m": n_mon, "s": n_sub, "i": n_igv, "t": n_tot,
                     "o": n_obs.strip() or None, "id": fid})
                st.session_state.pop("fa_editando", None)
                listo(f"Factura {n_comp.strip().upper()} actualizada." + (" Como cambió, vuelve a aprobación." if vuelve else ""))


def actualizar_ids(ids, sql, extra, mensaje):
    with db.motor().begin() as con:
        con.execute(text(sql), [{"id": int(i), **extra} for i in ids])
    listo(mensaje)


def aprobar(ids):
    n = f"{len(ids)} {'factura' if len(ids) == 1 else 'facturas'}"
    actualizar_ids(ids, "UPDATE facturas_proveedores SET estado = 'APROBADA', aprobado_por = :q, fecha_aprobacion = :f, "
                        "motivo_observacion = NULL WHERE id = :id AND estado = 'PENDIENTE'",
                   {"q": nombre_actual(), "f": hoy.isoformat()}, f"{n} {'aprobada' if len(ids) == 1 else 'aprobadas'}.")


def observar(ids, motivo):
    n = f"{len(ids)} {'factura' if len(ids) == 1 else 'facturas'}"
    actualizar_ids(ids, "UPDATE facturas_proveedores SET estado = 'OBSERVADA', aprobado_por = :q, fecha_aprobacion = :f, "
                        "motivo_observacion = :m WHERE id = :id AND estado = 'PENDIENTE'",
                   {"q": nombre_actual(), "f": hoy.isoformat(), "m": motivo},
                   f"{n} {'observada' if len(ids) == 1 else 'observadas'}. El administrador verá el motivo.")


REENVIAR = ("UPDATE facturas_proveedores SET estado = 'PENDIENTE', fecha_pago = NULL, aprobado_por = NULL, "
            "fecha_aprobacion = NULL, motivo_observacion = NULL WHERE id = :id")


def eliminar(ids):
    with db.motor().begin() as con:
        for t in ("facturas_items", "facturas_archivos"):
            con.execute(text(f"DELETE FROM {t} WHERE factura_id = :id"), [{"id": int(i)} for i in ids])
    actualizar_ids(ids, "DELETE FROM facturas_proveedores WHERE id = :id", {},
                   f"{len(ids)} {'factura eliminada' if len(ids) == 1 else 'facturas eliminadas'}.")


def releer(ids):
    with st.spinner("Volviendo a leer los documentos..."):
        releidas, sin_doc = releer_facturas(ids)
    aviso = f"Se {'volvió' if releidas == 1 else 'volvieron'} a leer {releidas} {'factura' if releidas == 1 else 'facturas'}."
    if sin_doc:
        aviso += f" {sin_doc} no {'tiene' if sin_doc == 1 else 'tienen'} documento guardado: adjúntalo con Editar."
    listo(aviso)


def ficha(f):
    """Ficha de una factura: documento, servicios, avance y las acciones que corresponden a su estado."""
    fid, estado = int(f["id"]), f["estado"]
    alerta_html = ""
    if f["alerta"]:
        alerta_html = estilo.estado_html(f["alerta"].replace("🔴 ", "").replace("🟡 ", ""), "no" if "🔴" in f["alerta"] else "espera")
    with st.container(border=True, key="ficha"):
        st.markdown(
            f'<div class="f-ficha"><h3>{html.escape(f["proveedor"] or "")}</h3>'
            f'<p class="sub">{html.escape(f["comprobante"] or "")}, emitida el {fecha_corta(f["fecha_emision"])}, '
            f'vence el {fecha_corta(f["fecha_vencimiento"])} {alerta_html}</p></div>', unsafe_allow_html=True)
        with st.expander("Ver documento"):
            ver_documento_guardado(fid, f"ficha_{fid}")
        items = cargar_items([fid])
        filas = "".join(f'<li><span>{html.escape(str(it.descripcion or "Servicio"))}</span>'
                        f'<span>{dinero(float(it.importe or 0), f["moneda"])}</span></li>' for it in items.itertuples())
        if not filas:
            filas = '<li><span style="color:var(--h-suave)">Sin detalle de servicios</span><span></span></li>'
        filas += (f'<li><span>Subtotal</span><span>{dinero(f["subtotal"], f["moneda"])}</span></li>'
                  f'<li><span>IGV</span><span>{dinero(f["igv"], f["moneda"])}</span></li>'
                  f'<li class="total"><span>Total</span><span>{dinero(f["total"], f["moneda"])}</span></li>')
        reg = fecha_corta((f["creado"] or "")[:10]) if isinstance(f["creado"], str) else ""
        if estado == "PENDIENTE":
            p2 = '<li class="ahora"><b>Por aprobar</b>Esperando al aprobador</li>'
        elif estado == "OBSERVADA":
            p2 = f'<li class="mal"><b>Observada</b>{html.escape(str(f["aprobado_por"] or ""))}</li>'
        elif estado == "ANULADA":
            p2 = '<li class="mal"><b>Anulada</b></li>'
        else:
            p2 = (f'<li class="hecho"><b>Aprobada</b>{html.escape(str(f["aprobado_por"] or ""))}'
                  f'{", " + fecha_corta(f["fecha_aprobacion"])[:5] if isinstance(f["fecha_aprobacion"], str) else ""}</li>')
        p3 = (f'<li class="hecho"><b>Pagada</b>{fecha_corta(f["fecha_pago"])}</li>' if estado == "PAGADA" else
              f'<li class="{"ahora" if estado == "APROBADA" else ""}"><b>Por pagar</b>'
              f'{"Vencida" if f["vencida"] else ""}</li>')
        st.markdown(f'<ul class="f-serv">{filas}</ul><ol class="f-pasos"><li class="hecho"><b>Registrada</b>{reg}</li>'
                    f'{p2}{p3}</ol>', unsafe_allow_html=True)
        if estado == "OBSERVADA" and f["motivo_observacion"]:
            st.warning(f"Observación de {f['aprobado_por']}: {f['motivo_observacion']}")

        if rol == "aprobador_facturas" or (puede_aprobar and estado == "PENDIENTE"):
            if estado == "PENDIENTE":
                motivo = st.text_input("Motivo (solo si vas a observar)", key=f"mot_{fid}")
                b1, b2 = st.columns(2)
                if b1.button("Aprobar", type="primary", key=f"apr_{fid}", use_container_width=True):
                    aprobar([fid])
                if b2.button("Observar", key=f"obs_{fid}", use_container_width=True):
                    if not motivo.strip():
                        st.error("Escribe el motivo para que el administrador pueda corregirla.")
                    else:
                        observar([fid], motivo.strip())
            if rol == "aprobador_facturas":
                return
            st.divider()

        if estado == "APROBADA":
            fp = st.date_input("Fecha de pago", value=hoy, format="DD/MM/YYYY", key=f"fp_{fid}")
            if st.button("Marcar como pagada", type="primary", key=f"pag_{fid}", use_container_width=True):
                actualizar_ids([fid], "UPDATE facturas_proveedores SET estado = 'PAGADA', fecha_pago = :f WHERE id = :id",
                               {"f": fp.isoformat()}, "Factura marcada como pagada.")
        b = st.columns(2) + st.columns(2)
        if b[0].button("Editar", key=f"edt_{fid}", use_container_width=True):
            st.session_state["fa_editando"] = fid
            st.rerun()
        if estado in ("OBSERVADA", "ANULADA", "PAGADA"):
            if b[1].button("Reenviar a aprobación" if estado != "PAGADA" else "Volver a por aprobar",
                           key=f"ree_{fid}", use_container_width=True):
                actualizar_ids([fid], REENVIAR, {}, "Factura enviada otra vez a aprobación.")
        else:
            if b[1].button("Volver a leer", key=f"rel_{fid}", use_container_width=True,
                           help="Relee el documento guardado con el lector actual"):
                releer([fid])
        if estado != "ANULADA" and b[2].button("Anular", key=f"anu_{fid}", use_container_width=True):
            actualizar_ids([fid], "UPDATE facturas_proveedores SET estado = 'ANULADA' WHERE id = :id", {}, "Factura anulada.")
        if st.button("Eliminar factura", type="tertiary", key=f"eli_{fid}"):
            st.session_state["fa_confirmar"] = [fid]
        if st.session_state.get("fa_confirmar") == [fid]:
            st.warning("¿Eliminar esta factura? No se puede deshacer.")
            if st.button("Sí, eliminar", key=f"eli2_{fid}"):
                st.session_state.pop("fa_confirmar")
                eliminar([fid])


# ---------- Vista principal: filtros, deuda, lista y ficha ----------
with tab_main:
    if facturas.empty:
        st.info("Todavía no hay facturas." + (" Cárgalas en la pestaña «Cargar facturas»." if rol == "admin" else ""))
    else:
        FILTROS = {"Pendientes de pago": NO_PAGADAS, "Por aprobar": ("PENDIENTE",), "Aprobadas": ("APROBADA",),
                   "Observadas": ("OBSERVADA",), "Pagadas": ("PAGADA",), "Anuladas": ("ANULADA",), "Todas": None}
        defecto = "Por aprobar" if rol == "aprobador_facturas" else "Pendientes de pago"
        filtro = st.pills("Estado", list(FILTROS), default=defecto, key="fa_filtro", label_visibility="collapsed") or defecto
        f1, f2, f3 = st.columns([2, 2, 1])
        proveedor = f1.selectbox("Proveedor", ["Todos los proveedores"] + sorted(facturas["proveedor"].dropna().unique()),
                                 key="fa_prov", label_visibility="collapsed")
        buscar = f2.text_input("Buscar", placeholder="Buscar proveedor, RUC o comprobante", key="fa_buscar",
                               label_visibility="collapsed")
        monedas = sorted(facturas["moneda"].unique())
        moneda = f3.selectbox("Moneda", monedas, key="fa_moneda", label_visibility="collapsed") if len(monedas) > 1 else monedas[0]

        base = facturas[facturas["moneda"] == moneda]
        if proveedor != "Todos los proveedores":
            base = base[base["proveedor"] == proveedor]
        if buscar:
            b_ = buscar.strip().lower()
            base = base[base[["proveedor", "ruc", "comprobante"]].fillna("").apply(lambda r: b_ in " ".join(r).lower(), axis=1)]
        vista = base if FILTROS[filtro] is None else base[base["estado"].isin(FILTROS[filtro])]

        # Lo que se debe: una sola barra (vencido, vence en 7 días, al día)
        pend = base[base["estado"].isin(NO_PAGADAS)]
        venc = pend[pend["vencida"]]
        prox = pend[(pend["dias"] >= 0) & (pend["dias"] <= DIAS_ALERTA)]
        total = float(pend["total"].sum())
        v_, p_ = float(venc["total"].sum()), float(prox["total"].sum())
        a_ = max(total - v_ - p_, 0)
        pct = (lambda x: 0 if total == 0 else max(x / total * 100, 1.5 if x else 0))
        sim = SIMBOLO.get(moneda, moneda)
        st.markdown(f"""<div class="f-deuda"><div class="f-deuda-cab"><b>{sim} {total:,.2f}</b>
<span>pendiente de pago en {len(pend)} {"factura" if len(pend) == 1 else "facturas"}</span></div>
<div class="f-barra"><i style="width:{pct(v_)}%;background:var(--h-rojo)"></i><i style="width:{pct(p_)}%;background:var(--h-ambar)"></i>
<i style="width:{pct(a_)}%;background:var(--h-azul)"></i></div>
<div class="f-ley"><div class="v">Vencido<strong>{sim} {v_:,.2f}</strong></div>
<div class="p">Vence en {DIAS_ALERTA} días<strong>{sim} {p_:,.2f}</strong></div>
<div class="a">Al día<strong>{sim} {a_:,.2f}</strong></div></div></div>""", unsafe_allow_html=True)

        izq, der = st.columns([1.4, 1], gap="medium")
        with izq:
            lista = vista.sort_values(["vencida", "fecha_vencimiento"], ascending=[False, True]).reset_index(drop=True)
            icono = lista["alerta"].map(lambda x: "🔴 " if "🔴" in str(x) else "🟡 " if "🟡" in str(x) else "")
            tabla = pd.DataFrame({
                "Proveedor": lista["proveedor"],
                "Vence": icono + lista["fecha_vencimiento"].map(fecha_corta),
                "Total": [dinero(t, m) for t, m in zip(lista["total"], lista["moneda"])], "Estado": lista["situacion"],
            })
            evento = st.dataframe(
                tabla, key=f"fa_lista_{st.session_state['fa_ronda']}", hide_index=True, width="stretch",
                on_select="rerun", selection_mode="multi-row", height=min(38 + 35 * max(len(tabla), 1), 460),
                column_config={"Proveedor": st.column_config.TextColumn(width=165),
                               "Vence": st.column_config.TextColumn(width=112, help="🔴 vencida, 🟡 vence en 7 días"),
                               "Total": st.column_config.TextColumn(width=98),
                               "Estado": st.column_config.TextColumn(width=96)})
            ids = [int(lista.iloc[k]["id"]) for k in evento.selection.rows]
            st.caption(f"{len(vista)} {'factura' if len(vista) == 1 else 'facturas'} en «{filtro}». "
                       "Marca una para ver su ficha, o varias para actuar sobre todas.")
            st.download_button("Descargar en Excel (lo filtrado)", data=excel(vista),
                               file_name=f"facturas_proveedores_{hoy:%Y-%m-%d}.xlsx",
                               mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                               disabled=vista.empty, key="fa_excel")

        with der:
            if len(ids) == 1:
                ficha(facturas[facturas["id"] == ids[0]].iloc[0])
            elif len(ids) > 1:
                sel = facturas[facturas["id"].isin(ids)]
                n = f"{len(ids)} facturas"
                with st.container(border=True, key="ficha"):
                    st.markdown(f"**{n} marcadas**, por {totales_por_moneda(sel)}.")
                    if rol == "aprobador_facturas":
                        motivo = st.text_input("Motivo (solo si vas a observar)", key="mot_varias")
                        b1, b2 = st.columns(2)
                        if b1.button(f"Aprobar {len(ids)}", type="primary", use_container_width=True):
                            aprobar(ids)
                        if b2.button(f"Observar {len(ids)}", use_container_width=True):
                            if motivo.strip():
                                observar(ids, motivo.strip())
                            else:
                                st.error("Escribe el motivo de la observación.")
                    else:
                        fp = st.date_input("Fecha de pago", value=hoy, format="DD/MM/YYYY", key="fp_varias")
                        b1, b2 = st.columns(2)
                        if b1.button("Marcar como pagadas", type="primary", use_container_width=True):
                            if set(sel["estado"]) != {"APROBADA"}:
                                st.error("Solo se pueden marcar como pagadas las facturas **aprobadas**.")
                            else:
                                actualizar_ids(ids, "UPDATE facturas_proveedores SET estado = 'PAGADA', fecha_pago = :f "
                                                    "WHERE id = :id", {"f": fp.isoformat()}, f"{n} marcadas como pagadas.")
                        if b2.button("Reenviar a aprobación", use_container_width=True):
                            actualizar_ids(ids, REENVIAR, {}, f"{n} enviadas otra vez a aprobación.")
                        b3, b4 = st.columns(2)
                        if b3.button("Volver a leer", use_container_width=True):
                            releer(ids)
                        if b4.button("Anular", use_container_width=True):
                            actualizar_ids(ids, "UPDATE facturas_proveedores SET estado = 'ANULADA' WHERE id = :id", {},
                                           f"{n} anuladas.")
            else:
                st.markdown('<div class="h-panel f-vacio">Elige una factura de la lista para ver su documento, '
                            'sus servicios y en qué paso está.</div>', unsafe_allow_html=True)

        editando = st.session_state.get("fa_editando")
        if rol == "admin" and editando and editando in set(facturas["id"]):
            editar_factura(facturas[facturas["id"] == editando].iloc[0])

        with st.expander("Ver gráficos"):
            g1, g2 = st.columns(2)
            with g1:
                st.markdown("**Monto por proveedor**")
                if vista.empty:
                    st.caption("Sin facturas en este filtro.")
                else:
                    por_prov = vista.groupby(["proveedor", "situacion"], as_index=False)["total"].sum() \
                        .rename(columns={"proveedor": "Proveedor", "situacion": "Estado", "total": "Monto"})
                    st.bar_chart(por_prov, x="Proveedor", y="Monto", color="Estado", horizontal=True,
                                 height=max(220, 38 * por_prov["Proveedor"].nunique()))
            with g2:
                st.markdown("**Antigüedad de lo pendiente**")
                tramos = ["Vencida más de 60 días", "Vencida 31 a 60 días", "Vencida 1 a 30 días",
                          f"Vence en 0 a {DIAS_ALERTA} días", f"Vence en más de {DIAS_ALERTA} días"]

                def tramo(d):
                    if d is None or pd.isna(d):
                        return tramos[4]
                    return (tramos[0] if d < -60 else tramos[1] if d < -30 else tramos[2] if d < 0
                            else tramos[3] if d <= DIAS_ALERTA else tramos[4])

                if pend.empty:
                    st.caption("No hay facturas pendientes de pago.")
                else:
                    ant = pend.assign(Tramo=pend["dias"].map(tramo)).groupby("Tramo")["total"].sum() \
                        .reindex(tramos, fill_value=0).reset_index().rename(columns={"total": "Monto"})
                    st.bar_chart(ant, x="Tramo", y="Monto", horizontal=True, height=240, color="#0B4F9C", sort=False)

if rol != "admin":
    st.stop()

# ---------- Cargar facturas (lectura automática) ----------
with tab_cargar:
    st.markdown("Sube una o varias facturas en **PDF** o **XML** (también el ZIP que manda el proveedor). "
                "El sistema busca la tabla de cada factura, sea cual sea el formato del proveedor, y saca los montos "
                "y el detalle de servicios. **Revisa antes de guardar**: con el XML la lectura es exacta; con el PDF "
                "puede haber algo que corregir.")
    if clave_api():
        st.caption("Lectura con IA activada para PDF y fotos.")
    if "fa_subida" not in st.session_state:
        st.session_state["fa_subida"] = 0
    ronda = st.session_state["fa_subida"]
    archivos = st.file_uploader("Facturas", type=["pdf", "xml", "zip", "png", "jpg", "jpeg", "webp"],
                                accept_multiple_files=True, key=f"fa_archivos_{ronda}")
    revisadas = []
    for k, a in enumerate(archivos or []):
        with st.spinner(f"Leyendo {a.name}..."):
            datos, nota = leer(a.name, a.getvalue(), VERSION)
        with st.container(border=True):
            st.markdown(f"**{a.name}**  \n:gray[{nota}]")
            with st.expander("Ver documento"):
                mostrar_documento(a.name, a.getvalue(), f"carga_{ronda}_{k}")
            c1, c2, c3 = st.columns([3, 2, 2])
            proveedor = c1.text_input("Proveedor", datos["proveedor"] or "", key=f"fa_p_{ronda}_{k}")
            ruc = c2.text_input("RUC", datos["ruc"] or "", key=f"fa_r_{ronda}_{k}")
            comprobante = c3.text_input("Comprobante", datos["comprobante"] or "", key=f"fa_c_{ronda}_{k}")
            c4, c5, c6 = st.columns(3)
            emision = c4.date_input("Emisión", date.fromisoformat(datos["fecha_emision"]) if datos["fecha_emision"] else None,
                                    format="DD/MM/YYYY", key=f"fa_e_{ronda}_{k}")
            vence = c5.date_input("Vence", date.fromisoformat(datos["fecha_vencimiento"]) if datos["fecha_vencimiento"] else None,
                                  format="DD/MM/YYYY", key=f"fa_v_{ronda}_{k}",
                                  help="Si queda vacío, se usa la fecha de emisión (pago al contado)")
            moneda = c6.selectbox("Moneda", ["PEN", "USD"], index=1 if datos["moneda"] == "USD" else 0,
                                  key=f"fa_m_{ronda}_{k}")
            c7, c8, c9 = st.columns(3)
            subtotal = c7.number_input("Subtotal", value=datos["subtotal"], min_value=0.0, format="%.2f",
                                       key=f"fa_s_{ronda}_{k}")
            igv = c8.number_input("IGV", value=datos["igv"], min_value=0.0, format="%.2f", key=f"fa_i_{ronda}_{k}")
            total = c9.number_input("Total", value=datos["total"], min_value=0.0, format="%.2f", key=f"fa_t_{ronda}_{k}")
            st.markdown("Detalle de servicios")
            items_df = st.data_editor(tabla_items(datos["items"]), key=f"fa_it_{ronda}_{k}", num_rows="dynamic",
                                      hide_index=True, width="stretch", column_config=COLS_ITEMS)
            items = limpiar_items(items_df)
            suma = round(sum(float(it["importe"] or 0) for it in items), 2)
            if items and subtotal and abs(suma - subtotal) > 0.05 and abs(suma - (total or 0)) > 0.05:
                st.caption(f":orange[La suma del detalle ({suma:,.2f}) no coincide con el subtotal ({subtotal:,.2f}) "
                           f"ni con el total. Revísalo.]")
            faltan = [n for n, v in (("proveedor", proveedor.strip()), ("comprobante", comprobante.strip()),
                                     ("emisión", emision), ("total", total)) if not v]
            if faltan:
                st.caption(":red[Falta completar: " + ", ".join(faltan) + "]")
            revisadas.append({"faltan": faltan, "fila": {
                "ruc": ruc.strip() or None, "proveedor": proveedor.strip(), "comprobante": comprobante.strip().upper(),
                "fecha_emision": emision.isoformat() if emision else None,
                "fecha_vencimiento": (vence or emision).isoformat() if (vence or emision) else None,
                "moneda": moneda, "subtotal": subtotal, "igv": igv, "total": total,
                "observacion": None, "archivo": a.name, "items": items, "documento": (a.name, a.getvalue()),
            }})

    if revisadas:
        incompletas = sum(bool(r["faltan"]) for r in revisadas)
        if incompletas:
            st.warning(f"{incompletas} {'factura tiene' if incompletas == 1 else 'facturas tienen'} datos por completar.")
        if st.button(f"Guardar {len(revisadas)} {'factura' if len(revisadas) == 1 else 'facturas'}",
                     type="primary", disabled=bool(incompletas)):
            guardadas, repetidas = guardar([r["fila"] for r in revisadas])
            aviso = f"Listo: se {'guardó' if guardadas == 1 else 'guardaron'} {guardadas} " \
                    f"{'factura' if guardadas == 1 else 'facturas'}."
            if repetidas:
                aviso += f" Se omitieron por estar ya registradas: {', '.join(map(str, repetidas))}."
            st.session_state["fa_aviso"] = aviso
            st.session_state["fa_subida"] += 1  # limpia la carga
            st.rerun()

# ---------- Agregar a mano ----------
with tab_mano:
    with st.form("fa_form", clear_on_submit=True):
        c1, c2 = st.columns(2)
        proveedor = c1.text_input("Proveedor")
        ruc = c2.text_input("RUC", max_chars=11)
        c3, c4, c5 = st.columns(3)
        comprobante = c3.text_input("Comprobante", placeholder="F001-00000123")
        emision = c4.date_input("Fecha de emisión", value=hoy, format="DD/MM/YYYY")
        vence = c5.date_input("Fecha de vencimiento", value=hoy + timedelta(days=30), format="DD/MM/YYYY")
        c6, c7, c8, c9 = st.columns(4)
        moneda = c6.selectbox("Moneda", ["PEN", "USD"])
        subtotal = c7.number_input("Subtotal", min_value=0.0, step=10.0, format="%.2f")
        igv = c8.number_input("IGV", min_value=0.0, step=1.0, format="%.2f")
        total = c9.number_input("Total", min_value=0.0, step=10.0, format="%.2f",
                                help="Si lo dejas en 0, se calcula como subtotal + IGV")
        observacion = st.text_input("Observación (opcional)")
        doc_mano = st.file_uploader("Documento de la factura (opcional)",
                                    type=["pdf", "xml", "zip", "png", "jpg", "jpeg", "webp"], key="fa_doc_mano")
        st.markdown("Detalle de servicios (opcional)")
        items_m = st.data_editor(tabla_items([]), key="fa_items_mano", num_rows="dynamic", hide_index=True,
                                 width="stretch", column_config=COLS_ITEMS)
        if st.form_submit_button("Guardar factura", type="primary"):
            total = total or round(subtotal + igv, 2)
            if not proveedor.strip() or not comprobante.strip() or total <= 0:
                st.error("Completa proveedor, comprobante y total.")
            elif vence < emision:
                st.error("La fecha de vencimiento es anterior a la de emisión.")
            else:
                guardadas, repetidas = guardar([{
                    "ruc": ruc.strip() or None, "proveedor": proveedor.strip(),
                    "comprobante": comprobante.strip().upper(), "fecha_emision": emision.isoformat(),
                    "fecha_vencimiento": vence.isoformat(), "moneda": moneda, "subtotal": subtotal or None,
                    "igv": igv or None, "total": total, "observacion": observacion.strip() or None, "archivo": None,
                    "items": limpiar_items(items_m),
                    "documento": (doc_mano.name, doc_mano.getvalue()) if doc_mano is not None else None,
                }])
                if repetidas:
                    st.error("Esa factura ya está registrada (mismo RUC y comprobante).")
                else:
                    st.session_state["fa_aviso"] = f"Factura {comprobante.strip().upper()} guardada."
                    st.rerun()
